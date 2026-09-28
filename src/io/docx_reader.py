from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.oxml.ns import qn

from src.domain.models import CellBorder, CellPresentation, ExtractedTable, StatementType
from src.extraction.statement_classifier import classify_table, get_title_hint
from src.normalization.text_cleaner import clean_text


def read_docx_tables(path: str | Path) -> list[ExtractedTable]:
    document = Document(str(path))
    tables: list[ExtractedTable] = []
    last_paragraph = ""
    recent_paragraphs: list[str] = []
    numbered_context: dict[int, str] = {}
    active_numbering_id: int | None = None
    last_numbered_paragraph = ""

    for block in _iter_document_blocks(document):
        if isinstance(block, Paragraph):
            text = clean_text(block.text)
            if text:
                last_paragraph = text
                recent_paragraphs.append(text)
                recent_paragraphs = recent_paragraphs[-8:]
                numbering = _paragraph_numbering(block)
                if numbering is not None:
                    numbering_id, level = numbering
                    last_numbered_paragraph = text
                    if numbering_id != active_numbering_id:
                        numbered_context.clear()
                        active_numbering_id = numbering_id
                    numbered_context = {
                        context_level: context
                        for context_level, context in numbered_context.items()
                        if context_level < level
                    }
                    numbered_context[level] = text
            continue

        table = block
        rows = [
            [clean_text(cell.text) for cell in row.cells]
            for row in table.rows
        ]
        fallback_title_hint = last_paragraph or get_title_hint(rows)
        context_hints = tuple(
            dict.fromkeys([*numbered_context.values(), *recent_paragraphs[-3:]])
        )
        statement_type = classify_table(
            rows,
            title_hint=fallback_title_hint,
            context_hints=context_hints,
        )
        title_hint = (
            last_numbered_paragraph
            if statement_type == StatementType.NOTE and last_numbered_paragraph
            else fallback_title_hint
        )
        tables.append(
            ExtractedTable(
                index=len(tables) + 1,
                rows=rows,
                statement_type=statement_type,
                title_hint=title_hint,
                context_hints=context_hints,
                cell_presentations=_table_cell_presentations(table),
                merged_ranges=_table_merged_ranges(table),
                column_widths=_table_column_widths(table),
            )
        )
        recent_paragraphs.clear()

    return tables


def read_docx_text_fragments(path: str | Path) -> tuple[str, ...]:
    """Đọc toàn bộ đoạn văn để nhận diện chính sách mà không đưa nội dung vào log."""
    document = Document(str(path))
    return tuple(
        text
        for paragraph in document.paragraphs
        if (text := clean_text(paragraph.text))
    )


def _iter_document_blocks(document):
    for child in document.element.body.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, document)
        elif isinstance(child, CT_Tbl):
            yield Table(child, document)


def _paragraph_numbering(paragraph: Paragraph) -> tuple[int, int] | None:
    properties = paragraph._p.pPr
    numbering = properties.numPr if properties is not None else None
    if numbering is None and paragraph.style is not None:
        style_properties = paragraph.style.element.pPr
        numbering = style_properties.numPr if style_properties is not None else None
    if numbering is None or numbering.numId is None:
        return None
    level = numbering.ilvl.val if numbering.ilvl is not None else 0
    return int(numbering.numId.val), int(level)


def _table_cell_presentations(table: Table) -> dict[tuple[int, int], CellPresentation]:
    presentations: dict[tuple[int, int], CellPresentation] = {}
    for row_idx, row in enumerate(table.rows):
        for col_idx, cell in enumerate(row.cells):
            presentations[(row_idx, col_idx)] = _cell_presentation(cell)
    return presentations


def _cell_presentation(cell) -> CellPresentation:
    runs = [run for paragraph in cell.paragraphs for run in paragraph.runs if clean_text(run.text)]
    bold_values = [_resolved_run_property(run, "bold") for run in runs]
    italic_values = [_resolved_run_property(run, "italic") for run in runs]
    underline_values = [_resolved_run_property(run, "underline") for run in runs]
    font_colors = [str(run.font.color.rgb) for run in runs if run.font.color.rgb is not None]

    tc_properties = cell._tc.get_or_add_tcPr()
    shading = tc_properties.find(qn("w:shd"))
    fill_color = _xml_color(shading.get(qn("w:fill"))) if shading is not None else ""
    borders = tc_properties.find(qn("w:tcBorders"))

    return CellPresentation(
        bold=_uniform_true(bold_values),
        italic=_uniform_true(italic_values),
        underline=_uniform_true(underline_values),
        font_color=_uniform_value(font_colors),
        fill_color=fill_color,
        horizontal_alignment=_horizontal_alignment(cell),
        vertical_alignment=_vertical_alignment(cell),
        top_border=_cell_border(borders, "top"),
        right_border=_cell_border(borders, "right"),
        bottom_border=_cell_border(borders, "bottom"),
        left_border=_cell_border(borders, "left"),
    )


def _resolved_run_property(run, property_name: str) -> bool | None:
    value = getattr(run, property_name)
    if value is not None:
        return bool(value)
    style_font = run._parent.style.font if run._parent.style is not None else None
    inherited = getattr(style_font, property_name, None) if style_font is not None else None
    return bool(inherited) if inherited is not None else None


def _uniform_true(values: list[bool | None]) -> bool:
    explicit = [value for value in values if value is not None]
    return bool(explicit) and all(explicit)


def _uniform_value(values: list[str]) -> str:
    return values[0] if values and all(value == values[0] for value in values) else ""


def _horizontal_alignment(cell) -> str:
    alignments = [paragraph.alignment for paragraph in cell.paragraphs if paragraph.alignment is not None]
    if not alignments or any(alignment != alignments[0] for alignment in alignments):
        return ""
    return {
        WD_ALIGN_PARAGRAPH.LEFT: "left",
        WD_ALIGN_PARAGRAPH.CENTER: "center",
        WD_ALIGN_PARAGRAPH.RIGHT: "right",
        WD_ALIGN_PARAGRAPH.JUSTIFY: "justify",
    }.get(alignments[0], "")


def _vertical_alignment(cell) -> str:
    return {
        WD_CELL_VERTICAL_ALIGNMENT.TOP: "top",
        WD_CELL_VERTICAL_ALIGNMENT.CENTER: "center",
        WD_CELL_VERTICAL_ALIGNMENT.BOTTOM: "bottom",
    }.get(cell.vertical_alignment, "")


def _cell_border(borders, side_name: str) -> CellBorder:
    if borders is None:
        return CellBorder()
    side = borders.find(qn(f"w:{side_name}"))
    if side is None:
        return CellBorder()
    word_style = side.get(qn("w:val"), "")
    style = {
        "single": "thin",
        "thick": "thick",
        "double": "double",
        "dashed": "dashed",
        "dashSmallGap": "dashed",
        "dotted": "dotted",
        "dotDash": "dashDot",
        "dotDotDash": "dashDotDot",
    }.get(word_style, "")
    return CellBorder(style=style, color=_xml_color(side.get(qn("w:color"))))


def _xml_color(value: str | None) -> str:
    if not value or value.lower() in {"auto", "none"}:
        return ""
    normalized = value.strip().lstrip("#").upper()
    return normalized if len(normalized) in {6, 8} else ""


def _table_merged_ranges(table: Table) -> tuple[tuple[int, int, int, int], ...]:
    coordinates_by_cell: dict[object, list[tuple[int, int]]] = {}
    for row_idx, row in enumerate(table.rows):
        for col_idx, cell in enumerate(row.cells):
            coordinates_by_cell.setdefault(cell._tc, []).append((row_idx, col_idx))

    ranges: list[tuple[int, int, int, int]] = []
    for coordinates in coordinates_by_cell.values():
        unique_coordinates = set(coordinates)
        if len(unique_coordinates) <= 1:
            continue
        min_row = min(row for row, _ in unique_coordinates)
        max_row = max(row for row, _ in unique_coordinates)
        min_col = min(col for _, col in unique_coordinates)
        max_col = max(col for _, col in unique_coordinates)
        if len(unique_coordinates) == (max_row - min_row + 1) * (max_col - min_col + 1):
            ranges.append((min_row, min_col, max_row, max_col))
    return tuple(sorted(ranges))


def _table_column_widths(table: Table) -> tuple[float, ...]:
    grid = table._tbl.tblGrid
    if grid is None:
        return ()
    widths: list[float] = []
    for grid_column in grid.gridCol_lst:
        width = grid_column.w
        widths.append(float(width.inches) if width is not None else 0.0)
    return tuple(widths)
