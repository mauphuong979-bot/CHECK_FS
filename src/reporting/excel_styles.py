from __future__ import annotations

from openpyxl.styles import Border, PatternFill, Side


COLOR_NAVY = "1F4E78"
COLOR_BLUE = "4472C4"
COLOR_LIGHT_BLUE = "D9EAF7"
COLOR_YELLOW = "FFF2CC"
COLOR_GRAY = "F2F2F2"
COLOR_BORDER = "D9D9D9"
COLOR_TEXT = "1F1F1F"
COLOR_MUTED_TEXT = "666666"
COLOR_HYPERLINK = "0563C1"
COLOR_WHITE = "FFFFFF"

TITLE_FILL = PatternFill("solid", fgColor=COLOR_NAVY)
HEADER_FILL = PatternFill("solid", fgColor=COLOR_BLUE)
GROUP_HEADER_FILL = PatternFill("solid", fgColor=COLOR_LIGHT_BLUE)
LINKED_VALUE_FILL = PatternFill("solid", fgColor=COLOR_GRAY)
WARNING_FILL = PatternFill("solid", fgColor=COLOR_YELLOW)
CHECK_VALUE_FILL = PatternFill("solid", fgColor=COLOR_GRAY)

STATUS_FILLS = {
    "Matched": CHECK_VALUE_FILL,
    "Difference": WARNING_FILL,
    "Note value not found": WARNING_FILL,
    "Statement not found": WARNING_FILL,
}

TABLE_STATUS_FILLS = {
    "Có rule nghiệp vụ": CHECK_VALUE_FILL,
    "Có kiểm tra": CHECK_VALUE_FILL,
    "Cần xem xét": WARNING_FILL,
    "Không có rule nghiệp vụ": CHECK_VALUE_FILL,
    "Không kiểm tra": CHECK_VALUE_FILL,
}

THIN_BORDER = Border(
    left=Side(style="thin", color=COLOR_BORDER),
    right=Side(style="thin", color=COLOR_BORDER),
    top=Side(style="thin", color=COLOR_BORDER),
    bottom=Side(style="thin", color=COLOR_BORDER),
)
