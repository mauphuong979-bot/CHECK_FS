from __future__ import annotations

from collections import defaultdict
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.shared import RGBColor

from src.domain.models import WordCommentCategory, WordCommentTarget


COMMENT_AUTHOR = "CHECK_FS"
COMMENT_INITIALS = "CFS"
COMMENT_DIFFERENCE_COLOR = RGBColor(0xC0, 0x00, 0x00)
COMMENT_REVIEW_COLOR = RGBColor(0xC6, 0x59, 0x11)
COMMENT_TRACE_COLOR = RGBColor(0x05, 0x63, 0xC1)


def build_commented_docx(
    source_path: str | Path,
    targets: list[WordCommentTarget],
) -> bytes:
    document = Document(str(source_path))
    if not hasattr(document, "add_comment"):
        raise RuntimeError("Phiên bản python-docx hiện tại chưa hỗ trợ tạo comment.")

    messages_by_anchor: dict[object, list[WordCommentTarget]] = defaultdict(list)
    cell_by_anchor: dict[object, object] = {}
    for target in targets:
        if not 1 <= target.table_index <= len(document.tables):
            continue
        table = document.tables[target.table_index - 1]
        if not 0 <= target.row_index < len(table.rows):
            continue
        if not 0 <= target.column_index < len(table.rows[target.row_index].cells):
            continue
        cell = _resolve_comment_cell(table, target.row_index, target.column_index)
        if cell is None:
            continue
        anchor = cell._tc
        cell_by_anchor[anchor] = cell
        if not any(item.message == target.message for item in messages_by_anchor[anchor]):
            messages_by_anchor[anchor].append(target)

    for anchor, targets_for_anchor in messages_by_anchor.items():
        runs = _nonempty_runs(cell_by_anchor[anchor])
        if not runs:
            continue
        comment = document.add_comment(
            runs=runs,
            text="",
            author=COMMENT_AUTHOR,
            initials=COMMENT_INITIALS,
        )
        _write_styled_comment(comment, targets_for_anchor)

    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def _write_styled_comment(comment, targets: list[WordCommentTarget]) -> None:
    paragraph = comment.paragraphs[0]
    for run in paragraph.runs:
        paragraph._p.remove(run._r)
    for index, target in enumerate(targets):
        if index:
            comment.add_paragraph()
            paragraph = comment.add_paragraph()
        _write_comment_message(paragraph, target)


def _write_comment_message(paragraph, target: WordCommentTarget) -> None:
    message = target.message
    trace_end = message.find(":") + 1 if message.startswith("Đối chiếu ") else 0
    if trace_end:
        trace_run = paragraph.add_run(message[:trace_end])
        trace_run.bold = True
        trace_run.font.color.rgb = COMMENT_TRACE_COLOR
        message = message[trace_end:]

    run = paragraph.add_run(message)
    if target.category == WordCommentCategory.DIFFERENCE:
        run.bold = True
        run.font.color.rgb = COMMENT_DIFFERENCE_COLOR
    elif target.category == WordCommentCategory.REVIEW:
        run.font.color.rgb = COMMENT_REVIEW_COLOR


def _resolve_comment_cell(table, row_index: int, column_index: int):
    preferred = table.rows[row_index].cells[column_index]
    if _nonempty_runs(preferred):
        return preferred
    for cell in table.rows[row_index].cells:
        if _nonempty_runs(cell):
            return cell
    for row in table.rows:
        for cell in row.cells:
            if _nonempty_runs(cell):
                return cell
    return None


def _nonempty_runs(cell) -> list:
    return [
        run
        for paragraph in cell.paragraphs
        for run in paragraph.runs
        if run.text.strip()
    ]
