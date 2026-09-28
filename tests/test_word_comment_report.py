from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from zipfile import ZipFile

from docx import Document

from src.domain.models import WordCommentCategory, WordCommentTarget
from src.reporting.word_comment_report import build_commented_docx


class WordCommentReportTest(unittest.TestCase):
    def test_adds_true_word_comment_and_aggregates_same_cell(self):
        with TemporaryDirectory() as temp_dir:
            source_path = Path(temp_dir) / "source.docx"
            document = Document()
            table = document.add_table(rows=1, cols=2)
            table.cell(0, 0).text = "Chỉ tiêu"
            table.cell(0, 1).text = "100"
            document.save(source_path)

            output = build_commented_docx(
                source_path,
                [
                    WordCommentTarget(1, 0, 1, "Có chênh lệch."),
                    WordCommentTarget(1, 0, 1, "Cần kiểm tra."),
                ],
            )

        reopened = Document(BytesIO(output))
        comments = list(reopened.comments)
        self.assertEqual(len(comments), 1)
        self.assertIn("Có chênh lệch.", comments[0].text)
        self.assertIn("Cần kiểm tra.", comments[0].text)
        with ZipFile(BytesIO(output)) as archive:
            names = set(archive.namelist())
            document_xml = archive.read("word/document.xml")
        self.assertIn("word/comments.xml", names)
        self.assertIn(b"commentRangeStart", document_xml)
        self.assertIn(b"commentRangeEnd", document_xml)
        self.assertIn(b"commentReference", document_xml)

    def test_empty_target_falls_back_without_adding_visible_text(self):
        with TemporaryDirectory() as temp_dir:
            source_path = Path(temp_dir) / "source.docx"
            document = Document()
            table = document.add_table(rows=1, cols=2)
            table.cell(0, 0).text = "Nhãn"
            table.cell(0, 1).text = ""
            document.save(source_path)

            output = build_commented_docx(
                source_path,
                [WordCommentTarget(1, 0, 1, "Không tìm thấy.")],
            )

        reopened = Document(BytesIO(output))
        self.assertEqual(reopened.tables[0].cell(0, 1).text, "")
        comments = list(reopened.comments)
        self.assertEqual(len(comments), 1)
        self.assertEqual(comments[0].text, "Không tìm thấy.")

    def test_formats_difference_and_review_comments_by_category(self):
        with TemporaryDirectory() as temp_dir:
            source_path = Path(temp_dir) / "source.docx"
            document = Document()
            table = document.add_table(rows=1, cols=1)
            table.cell(0, 0).text = "100"
            document.save(source_path)

            output = build_commented_docx(
                source_path,
                [
                    WordCommentTarget(
                        1, 0, 0, "Đối chiếu TM–BS: Có chênh lệch 10 VND.",
                        WordCommentCategory.DIFFERENCE,
                    ),
                    WordCommentTarget(
                        1, 0, 0, "Cần kiểm toán viên xác minh.",
                        WordCommentCategory.REVIEW,
                    ),
                ],
            )

        comment = list(Document(BytesIO(output)).comments)[0]
        self.assertEqual(
            comment.text,
            "Đối chiếu TM–BS: Có chênh lệch 10 VND.\n\nCần kiểm toán viên xác minh.",
        )
        runs = [run for paragraph in comment.paragraphs for run in paragraph.runs]
        self.assertEqual(str(runs[0].font.color.rgb), "0563C1")
        self.assertTrue(runs[0].bold)
        self.assertEqual(str(runs[1].font.color.rgb), "C00000")
        self.assertTrue(runs[1].bold)
        self.assertEqual(str(runs[-1].font.color.rgb), "C65911")


if __name__ == "__main__":
    unittest.main()
