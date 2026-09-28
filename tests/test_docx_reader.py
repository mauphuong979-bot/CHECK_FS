from pathlib import Path
from tempfile import TemporaryDirectory
from io import BytesIO
import unittest

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches
from openpyxl import load_workbook

from src.domain.models import StatementType
from src.io.docx_reader import read_docx_tables
from src.reporting.excel_report import build_audit_workbook


class DocxReaderPresentationTest(unittest.TestCase):
    def test_note_uses_nearest_numbered_heading_for_type_and_excel_title(self):
        with TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "numbered-note.docx"
            document = Document()
            heading = document.add_paragraph("Giá vốn hàng bán")
            properties = heading._p.get_or_add_pPr()
            numbering = OxmlElement("w:numPr")
            level = OxmlElement("w:ilvl")
            level.set(qn("w:val"), "0")
            number_id = OxmlElement("w:numId")
            number_id.set(qn("w:val"), "1")
            numbering.extend((level, number_id))
            properties.append(numbering)
            document.add_paragraph("Đơn vị tính: VND")
            table = document.add_table(rows=3, cols=3)
            table.rows[0].cells[1].text = "Năm nay"
            table.rows[0].cells[2].text = "Năm trước"
            table.rows[1].cells[0].text = "Giá vốn của thành phẩm đã bán"
            table.rows[1].cells[1].text = "100"
            table.rows[1].cells[2].text = "90"
            table.rows[2].cells[0].text = "Tổng cộng"
            table.rows[2].cells[1].text = "100"
            table.rows[2].cells[2].text = "90"
            document.save(path)

            extracted = read_docx_tables(path)[0]

        self.assertEqual(extracted.statement_type, StatementType.NOTE)
        self.assertEqual(extracted.title_hint, "Giá vốn hàng bán")
        workbook = load_workbook(BytesIO(build_audit_workbook([extracted])))
        self.assertEqual(workbook["T001_TM"]["B1"].value, "Giá vốn hàng bán")
        self.assertEqual(workbook["00_Tong_hop"]["C2"].value, "Giá vốn hàng bán")

    def test_preserves_cell_presentation_merges_and_relative_widths(self):
        with TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "presentation.docx"
            document = Document()
            table = document.add_table(rows=2, cols=3)
            table.columns[0].width = Inches(2.5)
            table.columns[1].width = Inches(1.0)
            table.columns[2].width = Inches(1.5)

            merged = table.cell(0, 0).merge(table.cell(0, 1))
            paragraph = merged.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run("Tiêu đề")
            run.bold = True
            run.italic = True
            run.underline = True
            merged.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

            properties = merged._tc.get_or_add_tcPr()
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), "D9EAF7")
            properties.append(shading)
            borders = OxmlElement("w:tcBorders")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "double")
            bottom.set(qn("w:color"), "1F4E78")
            borders.append(bottom)
            properties.append(borders)

            table.cell(1, 0).text = "Tổng cộng"
            table.cell(1, 1).text = "100"
            table.cell(1, 2).text = "90"
            document.save(path)

            extracted = read_docx_tables(path)[0]

        presentation = extracted.cell_presentations[(0, 0)]
        self.assertTrue(presentation.bold)
        self.assertTrue(presentation.italic)
        self.assertTrue(presentation.underline)
        self.assertEqual(presentation.fill_color, "D9EAF7")
        self.assertEqual(presentation.horizontal_alignment, "center")
        self.assertEqual(presentation.vertical_alignment, "center")
        self.assertEqual(presentation.bottom_border.style, "double")
        self.assertIn((0, 0, 0, 1), extracted.merged_ranges)
        self.assertGreater(extracted.column_widths[0], extracted.column_widths[1])


if __name__ == "__main__":
    unittest.main()
