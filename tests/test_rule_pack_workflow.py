from io import BytesIO
import unittest
from unittest.mock import patch

from openpyxl import load_workbook

from src.domain.audit_rules import EMPTY_RULE_PACK, TT200_RULE_PACK
from src.domain.models import ExtractedTable, StatementType
from src.reporting.excel_report import build_audit_workbook


class RulePackWorkflowTest(unittest.TestCase):
    def test_empty_rule_pack_skips_tt200_note_reconciliation(self):
        table = ExtractedTable(
            1,
            [["Nội dung", "Cuối năm"], ["Tổng cộng", "100"]],
            StatementType.NOTE,
            "Thuyết minh khác",
        )
        with patch("src.reporting.excel_report.reconcile_note_tables") as reconcile:
            workbook_bytes = build_audit_workbook([table], rule_pack=EMPTY_RULE_PACK)

        reconcile.assert_not_called()
        workbook = load_workbook(BytesIO(workbook_bytes))
        self.assertEqual(workbook["00_Tong_hop"]["N1"].value, EMPTY_RULE_PACK.display_name)

    def test_tt200_rule_pack_is_recorded_in_summary(self):
        workbook = load_workbook(BytesIO(build_audit_workbook([], rule_pack=TT200_RULE_PACK)))

        summary = workbook["00_Tong_hop"]
        self.assertEqual(summary["N1"].value, "Thông tư 200/2014/TT-BTC")
        self.assertEqual(summary["N2"].value, "tt200-v1")


if __name__ == "__main__":
    unittest.main()
