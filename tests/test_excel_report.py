from io import BytesIO
import re
import unittest

from openpyxl import load_workbook
from openpyxl import Workbook

from src.domain.models import (
    CellBorder,
    CellPresentation,
    ExtractedTable,
    NoteMatchResult,
    StatementType,
    TableCheckResult,
    WordCommentCategory,
)
from src.reporting.excel_report import (
    COLOR_BLUE,
    COLOR_NAVY,
    NOTE_STATUS_LABELS,
    SUMMARY_SORT_ISSUES_DESC,
    SUMMARY_SORT_STATUS,
    WARNING_FILL,
    _note_reconciliation_summary_fill,
    _note_reconciliation_summary_status,
    _sort_summary_tables,
    _looks_like_total_label,
    _write_note_reconciliation_detail,
    build_audit_workbook,
    build_audit_workbook_with_comment_targets,
)


class ExcelReportTest(unittest.TestCase):
    def test_hai_luong_total_patterns_include_all_business_components(self):
        revenue = ExtractedTable(
            1,
            [
                ["", "Năm nay", "", "Năm trước"],
                ["", "VND", "", "VND"],
                ["", "", "", ""],
                ["Doanh thu bán thành phẩm", "1.000", "", "900"],
                ["", "1.000", "", "900"],
                ["Trừ:", "", "", ""],
                ["Các khoản giảm trừ doanh thu", "100", "", "50"],
                ["", "", "", ""],
                ["Doanh thu thuần", "900", "", "850"],
            ],
            StatementType.NOTE,
            "Doanh thu thuần về bán hàng và cung cấp dịch vụ",
        )
        other_expense = ExtractedTable(
            2,
            [
                ["Giá trị còn lại của tài sản cố định thanh lý", "7", "", "3"],
                ["Giá trị nguyên vật liệu báo phế", "5", "", "2"],
                ["Điều chỉnh công nợ", "1", "", "-"],
                ["Tổng cộng", "13", "", "5"],
            ],
            StatementType.NOTE,
            "Chi phí khác",
        )

        workbook = load_workbook(
            BytesIO(build_audit_workbook([revenue, other_expense])),
            data_only=False,
        )
        revenue_values = [cell.value for row in workbook["T001_TM"] for cell in row]
        expense_values = [cell.value for row in workbook["T002_TM"] for cell in row]

        self.assertIn("=ROUND(B12-B8+B10,2)", revenue_values)
        self.assertIn("=ROUND(D12-D8+D10,2)", revenue_values)
        self.assertIn("=ROUND(B7-SUM(B4:B6),2)", expense_values)
        self.assertIn("=ROUND(D7-SUM(D4:D6),2)", expense_values)

    def test_income_tax_expense_includes_global_minimum_tax(self):
        table = ExtractedTable(
            1,
            [
                ["Lợi nhuận trước thuế", "100"],
                ["Điều chỉnh tăng lợi nhuận trước thuế", "-"],
                ["Điều chỉnh giảm lợi nhuận trước thuế", "-"],
                ["Thu nhập chịu thuế", "100"],
                ["Số chuyển lỗ mang sang", "-"],
                ["Thu nhập tính thuế", "100"],
                ["Thuế thu nhập doanh nghiệp phải nộp ước tính", "20"],
                ["Thuế tối thiểu toàn cầu", "5"],
                ["Chi phí thuế thu nhập doanh nghiệp hiện hành", "25"],
            ],
            StatementType.NOTE,
            "Chi phí thuế thu nhập doanh nghiệp hiện hành",
        )

        workbook = load_workbook(BytesIO(build_audit_workbook([table])), data_only=False)
        values = [cell.value for row in workbook["T001_TM"] for cell in row]

        self.assertIn("=ROUND(B12-B10-B11,2)", values)

    def test_parent_detail_reports_not_applicable_when_maturity_split_follows(self):
        statement = ExtractedTable(
            1,
            [
                ["Chỉ tiêu", "Mã số", "Thuyết minh", "Năm nay", "Năm trước"],
                ["Phải thu ngắn hạn khác", "136", "V.3", "10", "-"],
                ["Phải thu dài hạn khác", "216", "V.3", "90", "80"],
            ],
            StatementType.BALANCE_SHEET_ASSETS,
        )
        parent = ExtractedTable(
            2,
            [["Khoản A", "90", "80"], ["Khoản B", "10", "-"], ["Tổng cộng", "100", "80"]],
            StatementType.NOTE,
            "3. Phải thu khác",
            context_hints=("3. Phải thu khác",),
        )
        split = ExtractedTable(
            3,
            [["Ngắn hạn", "10", "-"], ["Dài hạn", "90", "80"], ["Tổng cộng", "100", "80"]],
            StatementType.NOTE,
            "3. Phải thu khác",
            context_hints=("3. Phải thu khác",),
        )

        workbook_bytes, comment_targets = build_audit_workbook_with_comment_targets(
            [statement, parent, split]
        )
        workbook = load_workbook(BytesIO(workbook_bytes), data_only=False)

        self.assertEqual(workbook["00_Tong_hop"]["J3"].value, "Không áp dụng")
        self.assertEqual(workbook["00_Tong_hop"]["J4"].value, "Khớp")
        self.assertIn(
            "Không áp dụng đối chiếu trực tiếp; số liệu được đối chiếu theo bảng phân loại ngắn hạn/dài hạn bên dưới.",
            [cell.value for row in workbook["T002_TM"] for cell in row],
        )
        self.assertFalse(
            any(target.table_index == 2 for target in comment_targets)
        )

    def test_note_reconciliation_formulas_and_links_use_merged_anchor(self):
        table = ExtractedTable(
            1,
            [
                ["", "", "", "Tổng cộng", "Tổng cộng", ""],
                ["", "", "", "VND", "VND", ""],
                ["", "", "", "", "", ""],
                ["Số đầu năm", "", "", "90", "90", ""],
                ["Số cuối năm", "", "", "", "100", "100"],
            ],
            StatementType.NOTE,
            "Tài sản cố định hữu hình",
            merged_ranges=((3, 3, 3, 4), (4, 4, 4, 5)),
        )
        match = NoteMatchResult(
            note_table_index=1,
            note_title="Tài sản cố định hữu hình",
            item_name="Tài sản cố định hữu hình",
            status="Matched",
            note_current_cell="E8",
            note_prior_cell="E7",
            statement_current_cell="D5",
            statement_prior_cell="E5",
            statement_table_index=2,
        )
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "T001_TM"

        _write_note_reconciliation_detail(
            worksheet,
            [match],
            {2: "T002_BS_TS"},
            table,
        )

        formulas = {
            cell.value: cell.hyperlink.target if cell.hyperlink else None
            for row in worksheet.iter_rows()
            for cell in row
            if isinstance(cell.value, str) and cell.value.startswith("=ROUND(")
        }
        self.assertIn("=ROUND(N(E4)-N(E8),2)", formulas)
        self.assertIn("=ROUND(N(D4)-N(D7),2)", formulas)
        self.assertEqual(formulas["=ROUND(N(D4)-N(D7),2)"], "#T001_TM!D7")

    def test_horizontal_check_ignores_merged_duplicates_and_comments_anchor(self):
        table = ExtractedTable(
            1,
            [
                ["", "Nhà cửa", "Máy móc", "Tổng cộng", "Tổng cộng"],
                ["", "VND", "VND", "VND", "VND"],
                ["", "", "", "", ""],
                ["Số đầu năm", "40", "50", "100", "100"],
            ],
            StatementType.NOTE,
            "Tài sản cố định hữu hình",
            merged_ranges=(
                (0, 3, 0, 4),
                (1, 3, 1, 4),
                (3, 3, 3, 4),
            ),
        )

        workbook_bytes, _ = build_audit_workbook_with_comment_targets([table])
        worksheet = load_workbook(BytesIO(workbook_bytes), data_only=False)["T001_TM"]

        self.assertEqual(worksheet["F7"].value, "=ROUND((D7-SUM(B7,C7)),2)")
        self.assertIsNotNone(worksheet["D7"].comment)
        self.assertIn("chênh lệch 10", worksheet["D7"].comment.text)
        self.assertTrue(
            {"D4:E4", "D5:E5", "D7:E7"}.issubset(
                {str(cell_range) for cell_range in worksheet.merged_cells.ranges}
            )
        )

    def test_summary_reconciliation_status_uses_conservative_priority(self):
        note = ExtractedTable(1, [["Tổng cộng", "100"]], StatementType.NOTE)
        matches = [
            NoteMatchResult(1, "TM", status="Matched"),
            NoteMatchResult(1, "TM", status="Note value not found"),
            NoteMatchResult(1, "TM", status="Difference"),
        ]

        self.assertEqual(
            _note_reconciliation_summary_status(note, matches),
            "Có chênh lệch",
        )
        self.assertEqual(
            _note_reconciliation_summary_status(note, matches[:2]),
            "Không tìm thấy số liệu TM",
        )
        self.assertEqual(
            _note_reconciliation_summary_status(note, matches[:1]),
            "Khớp",
        )
        self.assertEqual(
            _note_reconciliation_summary_status(note, []),
            "Chưa có kết quả",
        )
        self.assertTrue(
            _note_reconciliation_summary_fill("Khớp").fgColor.rgb.endswith("F2F2F2")
        )
        self.assertTrue(
            _note_reconciliation_summary_fill("Có chênh lệch").fgColor.rgb.endswith(
                "FFF2CC"
            )
        )

    def test_receivable_net_continuation_links_previous_total_and_collapses_merged_duplicates(self):
        detail = ExtractedTable(
            1,
            [
                ["", "Số cuối năm", "", "Số đầu năm"],
                ["", "VND", "", "VND"],
                ["", "", "", ""],
                ["Tổng cộng", "100", "", "50"],
            ],
            StatementType.NOTE,
            "Phải thu ngắn hạn của khách hàng",
        )
        continuation = ExtractedTable(
            2,
            [
                ["Dự phòng phải thu ngắn hạn khó đòi (Thuyết minh V.3)", "(10)", "(10)", "(10)", "", "", "(5)", "(5)"],
                ["Giá trị thuần", "Giá trị thuần", "90", "", "", "45", "45"],
            ],
            StatementType.NOTE,
            "Phải thu ngắn hạn của khách hàng",
            merged_ranges=(
                (0, 1, 0, 3),
                (0, 6, 0, 7),
                (1, 0, 1, 1),
                (1, 5, 1, 6),
            ),
        )

        workbook = load_workbook(
            BytesIO(build_audit_workbook([detail, continuation])),
            data_only=False,
        )
        worksheet = workbook["T002_TM"]

        self.assertEqual(worksheet["I5"].value, "=ROUND(C5-'T001_TM'!B7-B4,2)")
        self.assertEqual(worksheet["J5"].value, "=ROUND(F5-'T001_TM'!D7-G4,2)")
        self.assertNotEqual(worksheet["C5"].fill.fgColor.rgb, WARNING_FILL.fgColor.rgb)
        self.assertNotEqual(worksheet["F5"].fill.fgColor.rgb, WARNING_FILL.fgColor.rgb)

    def test_income_tax_schedule_writes_structured_vertical_checks(self):
        table = ExtractedTable(
            1,
            [
                ["", "Năm nay", "", "Năm trước"],
                ["", "VND", "", "VND"],
                ["", "", "", ""],
                ["Lợi nhuận trước thuế", "(10)", "", "(180)"],
                ["Ảnh hưởng do:", "", "", ""],
                ["Điều chỉnh tăng lợi nhuận trước thuế", "45", "", "81"],
                ["Điều chỉnh giảm lợi nhuận trước thuế", "-", "", "1"],
                ["Thu nhập chịu thuế", "35", "", "(100)"],
                ["Trong đó:", "", "", ""],
                ["Thu nhập từ hoạt động sản xuất kinh doanh", "35", "", "(100)"],
                ["Các khoản thu nhập khác", "-", "", "-"],
                ["", "", "", ""],
                ["Số chuyển lỗ mang sang", "35", "", "-"],
                ["Thu nhập tính thuế", "-", "", "(100)"],
                ["Trong đó:", "", "", ""],
                ["Thu nhập từ hoạt động sản xuất kinh doanh", "-", "", "(100)"],
                ["Các khoản thu nhập khác", "-", "", "-"],
                ["", "", "", ""],
                ["Thuế thu nhập doanh nghiệp phải nộp ước tính từ hoạt động SXKD", "20", "", "30"],
                ["Thuế thu nhập doanh nghiệp phải nộp ước tính từ thu nhập khác", "-", "", "2"],
                ["Chi phí thuế thu nhập doanh nghiệp hiện hành", "20", "", "32"],
            ],
            StatementType.NOTE,
            "Chi phí thuế thu nhập doanh nghiệp hiện hành",
        )

        workbook = load_workbook(BytesIO(build_audit_workbook([table])), data_only=False)
        worksheet = workbook["T001_TM"]
        formulas = {
            cell.value
            for row in worksheet.iter_rows()
            for cell in row
            if isinstance(cell.value, str) and cell.value.startswith("=")
        }

        self.assertIn("=ROUND(B11-B7-B9+ABS(B10),2)", formulas)
        self.assertIn("=ROUND(D11-D7-D9+ABS(D10),2)", formulas)
        self.assertIn("=ROUND(B11-B13-B14,2)", formulas)
        self.assertIn("=ROUND(B17-B11+ABS(B16),2)", formulas)
        self.assertIn("=ROUND(B17-B19-B20,2)", formulas)
        self.assertIn("=ROUND(B24-B22-B23,2)", formulas)
        self.assertNotEqual(worksheet["B11"].fill.fgColor.rgb, WARNING_FILL.fgColor.rgb)
        self.assertNotEqual(worksheet["D11"].fill.fgColor.rgb, WARNING_FILL.fgColor.rgb)

    def test_income_tax_schedule_uses_only_business_rules_with_exemption(self):
        table = ExtractedTable(
            1,
            [
                ["", "Năm nay", "", "Năm trước"],
                ["", "VND", "", "VND"],
                ["", "", "", ""],
                ["Lợi nhuận trước thuế", "242.390.750.439", "", "(14.665.951.431)"],
                ["Ảnh hưởng do:", "", "", ""],
                ["Điều chỉnh tăng lợi nhuận trước thuế", "10.279.063.186", "", "2.148.100.291"],
                ["Điều chỉnh giảm lợi nhuận trước thuế", "(3.470.802.266)", "", "(1.886.762.280)"],
                ["Thu nhập chịu thuế", "249.199.011.359", "", "(14.404.613.420)"],
                ["Trong đó:", "", "", ""],
                ["Thu nhập từ hoạt động chế xuất", "249.055.127.093", "", "(14.404.613.420)"],
                ["Thu nhập từ hoạt động thương mại", "44.063.139", "", "-"],
                ["Các khoản thu nhập khác", "99.821.127", "", "-"],
                ["", "", "", ""],
                ["Số chuyển lỗ mang sang", "(19.701.639.869)", "", "-"],
                ["Thu nhập tính thuế", "229.497.371.490", "", "(14.404.613.420)"],
                ["Trong đó:", "", "", ""],
                ["Thu nhập từ hoạt động chế xuất", "229.353.487.224", "", "(14.404.613.420)"],
                ["Thu nhập từ hoạt động thương mại", "44.063.139", "", "-"],
                ["Các khoản thu nhập khác", "99.821.127", "", "-"],
                ["", "229.497.371.490", "", "(14.404.613.420)"],
                ["", "", "", ""],
                ["Thuế thu nhập doanh nghiệp phải nộp ước tính từ thu nhập từ hoạt động chế xuất", "45.870.697.445", "", "-"],
                ["Thuế thu nhập doanh nghiệp phải nộp ước tính từ thu nhập từ hoạt động thương mại", "8.812.628", "", "-"],
                ["Thuế thu nhập doanh nghiệp phải nộp ước tính từ các khoản thu nhập khác", "19.964.225", "", "-"],
                ["", "", "", ""],
                ["Thuế thu nhập doanh nghiệp được miễn trong năm", "(45.870.697.445)", "", "-"],
                ["", "", "", ""],
                ["Chi phí thuế thu nhập doanh nghiệp hiện hành", "28.776.853", "", "-"],
            ],
            StatementType.NOTE,
            "Chi phí thuế thu nhập doanh nghiệp hiện hành",
        )

        workbook = load_workbook(BytesIO(build_audit_workbook([table])), data_only=False)
        worksheet = workbook["T001_TM"]
        formulas = {
            cell.value
            for row in worksheet.iter_rows()
            for cell in row
            if isinstance(cell.value, str) and cell.value.startswith("=")
        }

        self.assertIn("=ROUND(B18-B11+ABS(B17),2)", formulas)
        self.assertIn("=ROUND(B31-B25-B26-B27+ABS(B29),2)", formulas)
        self.assertNotEqual(worksheet["B18"].fill.fgColor.rgb, WARNING_FILL.fgColor.rgb)
        self.assertNotEqual(worksheet["D18"].fill.fgColor.rgb, WARNING_FILL.fgColor.rgb)
        self.assertNotEqual(worksheet["B31"].fill.fgColor.rgb, WARNING_FILL.fgColor.rgb)

    def test_tax_rollforward_horizontal_check_accepts_positive_or_negative_paid_style(self):
        table = ExtractedTable(
            1,
            [
                ["", "Số đầu năm", "", "Số phải nộp", "", "Số đã nộp, khấu trừ", "", "Số cuối năm"],
                ["", "VND", "", "VND", "", "VND", "", "VND"],
                ["", "", "", "", "", "", "", ""],
                ["Thuế GTGT - trình bày âm", "100", "", "50", "", "(30)", "", "120"],
                ["Thuế GTGT - trình bày dương", "100", "", "50", "", "30", "", "120"],
                ["Thuế bằng không", "-", "", "-", "", "-", "", "-"],
                ["Thuế có chênh lệch", "100", "", "50", "", "30", "", "121"],
            ],
            StatementType.NOTE,
            "Thuế và các khoản phải nộp Nhà nước",
        )

        workbook = load_workbook(BytesIO(build_audit_workbook([table])), data_only=False)
        worksheet = workbook["T001_TM"]

        self.assertEqual(worksheet["I7"].value, "=ROUND(H7-B7-D7+ABS(F7),2)")
        self.assertEqual(worksheet["I8"].value, "=ROUND(H8-B8-D8+ABS(F8),2)")
        self.assertEqual(worksheet["I9"].value, "=ROUND(H9-B9-D9+ABS(F9),2)")
        self.assertEqual(worksheet["I10"].value, "=ROUND(H10-B10-D10+ABS(F10),2)")
        self.assertNotEqual(worksheet["H7"].fill.fgColor.rgb, WARNING_FILL.fgColor.rgb)
        self.assertNotEqual(worksheet["H8"].fill.fgColor.rgb, WARNING_FILL.fgColor.rgb)
        self.assertEqual(worksheet["H10"].fill.fgColor.rgb, WARNING_FILL.fgColor.rgb)

    def test_recognizes_specific_net_value_and_balance_total_labels(self):
        recognized_labels = (
            "gia tri thuan",
            "gia tri thuan cuoi nam",
            "so du cuoi nam",
            "so du cuoi ky",
            "so cuoi nam",
            "so cuoi ky",
        )

        for label in recognized_labels:
            with self.subTest(label=label):
                self.assertTrue(_looks_like_total_label(label))

        self.assertFalse(_looks_like_total_label("gia tri hang ton kho"))
        self.assertFalse(_looks_like_total_label("so du dau nam"))
        self.assertFalse(_looks_like_total_label("so du tai khoan chi tiet"))

    def test_net_value_and_closing_balance_rows_receive_vertical_checks(self):
        tables = [
            ExtractedTable(
                1,
                [
                    ["Nguyên liệu, vật liệu", "100", "90"],
                    ["Công cụ, dụng cụ", "50", "40"],
                    ["Dự phòng giảm giá", "-10", "-5"],
                    ["Giá trị thuần", "140", "125"],
                ],
                StatementType.NOTE,
                "Hàng tồn kho",
            ),
            ExtractedTable(
                2,
                [
                    ["Số dư đầu năm", "100", "80"],
                    ["Trích lập trong năm", "20", "15"],
                    ["Hoàn nhập trong năm", "-10", "-5"],
                    ["Số dư cuối năm", "110", "90"],
                ],
                StatementType.NOTE,
                "Tình hình tăng, giảm",
            ),
        ]

        workbook = load_workbook(BytesIO(build_audit_workbook(tables)), data_only=False)

        first_values = [cell.value for row in workbook["T001_TM"] for cell in row]
        second_values = [cell.value for row in workbook["T002_TM"] for cell in row]
        self.assertIn("Kiểm tra cộng dọc dòng 7", first_values)
        self.assertIn("=ROUND(B7-SUM(B4:B6),2)", first_values)
        self.assertIn("Kiểm tra cộng dọc dòng 7", second_values)
        self.assertIn("=ROUND(B7-SUM(B4:B6),2)", second_values)
        self.assertNotIn("Kiểm tra cộng dọc", first_values)
        self.assertNotIn("Kiểm tra cộng dọc", second_values)

    def test_closing_balance_includes_opening_balance_and_skips_empty_opening_check(self):
        table = ExtractedTable(
            1,
            [
                ["", "Năm nay", "Năm trước"],
                ["", "VND", "VND"],
                ["", "", ""],
                ["Số đầu năm", "(6.425.338.569)", "-"],
                ["Trích lập trong năm", "32.462.401.786", "100"],
                ["Hoàn nhập trong năm", "(33.501.873.655)", "-"],
                ["Số cuối năm", "(7.464.810.438)", "100"],
            ],
            StatementType.NOTE,
            "Tình hình tăng giảm dự phòng",
        )

        workbook = load_workbook(BytesIO(build_audit_workbook([table])), data_only=False)
        worksheet = workbook["T001_TM"]
        values = [cell.value for row in worksheet for cell in row]

        self.assertIn("Kiểm tra cộng dọc dòng 10", values)
        self.assertIn("=ROUND(B10-SUM(B7:B9),2)", values)
        self.assertNotIn("Kiểm tra cộng dọc dòng 7", values)

    def test_grand_total_crosses_blank_rows_and_group_headers(self):
        table = ExtractedTable(
            1,
            [
                ["", "Số cuối năm", "", "Số đầu năm"],
                ["", "VND", "", "VND"],
                ["", "", "", ""],
                ["Bên liên quan", "", "", ""],
                ["Công ty A", "28.145.742.850", "", "4.964.000.000"],
                ["Công ty B", "19.660.089.186", "", "10.439.141.928"],
                ["Công ty C", "4.562.175.062", "", "-"],
                ["Công ty D", "1.988.527.712", "", "-"],
                ["Công ty E", "290.758.550", "", "-"],
                ["", "", "", ""],
                ["Bên thứ ba", "", "", ""],
                ["Khách hàng A", "421.583.816.849", "", "20.000.000.000"],
                ["", "", "", ""],
                ["Khách hàng B", "145.250.212.664", "", "39.300.208.372"],
                ["Tổng cộng", "621.481.110.209", "", "74.703.350.300"],
            ],
            StatementType.NOTE,
            "Phải thu ngắn hạn của khách hàng",
        )

        workbook = load_workbook(BytesIO(build_audit_workbook([table])), data_only=False)
        values = [cell.value for row in workbook["T001_TM"] for cell in row]

        self.assertIn("=ROUND(B18-SUM(B8:B17),2)", values)
        self.assertIn("=ROUND(D18-SUM(D8:D17),2)", values)

    def test_closing_balance_includes_prior_closing_current_opening_transition(self):
        table = ExtractedTable(
            1,
            [
                ["", "Của chủ sở hữu", "Chưa phân phối", "Tổng cộng"],
                ["", "VND", "VND", "VND"],
                ["", "", "", ""],
                ["Số đầu năm trước", "529.677.417.162", "(82.425.728.587)", "447.251.688.575"],
                ["Lợi nhuận thuần trong năm trước", "-", "(180.809.995.791)", "(180.809.995.791)"],
                ["Số cuối năm trước, đầu năm nay", "529.677.417.162", "(263.235.724.378)", "266.441.692.784"],
                ["Tăng vốn trong năm", "357.980.000.000", "-", "357.980.000.000"],
                ["Lợi nhuận thuần trong năm", "-", "(10.662.321.436)", "(10.662.321.436)"],
                ["Số cuối năm", "887.657.417.162", "(273.898.045.814)", "613.759.371.348"],
            ],
            StatementType.NOTE,
            "Theo giấy chứng nhận đăng ký doanh nghiệp",
        )

        workbook = load_workbook(BytesIO(build_audit_workbook([table])), data_only=False)
        values = [cell.value for row in workbook["T001_TM"] for cell in row]

        self.assertIn("Kiểm tra cộng dọc dòng 12", values)
        self.assertIn("=ROUND(B12-SUM(B9:B11),2)", values)
        self.assertIn("=ROUND(D12-SUM(D9:D11),2)", values)

    def test_formatted_subtotals_and_net_revenue_receive_structured_checks(self):
        subtotal_presentation = CellPresentation(
            bold=True,
            bottom_border=CellBorder(style="double", color="1F1F1F"),
        )
        table = ExtractedTable(
            1,
            [
                ["", "Năm nay", "", "Năm trước"],
                ["", "VND", "", "VND"],
                ["", "", "", ""],
                ["Doanh thu bán thành phẩm", "1.304.311.708.114", "", "309.176.013.192"],
                ["Doanh thu bán phế liệu", "1.560.954.918", "", "415.661.523"],
                ["", "1.305.872.663.032", "", "309.591.674.715"],
                ["Trừ:", "", "", ""],
                ["Hàng bán bị trả lại", "1.336.300.686", "", "-"],
                ["", "1.336.300.686", "", "-"],
                ["", "", "", ""],
                ["Doanh thu thuần", "1.304.536.362.346", "", "309.591.674.715"],
            ],
            StatementType.NOTE,
            "Doanh thu thuần về bán hàng và cung cấp dịch vụ",
            cell_presentations={
                (5, 1): subtotal_presentation,
                (5, 3): subtotal_presentation,
                (8, 1): subtotal_presentation,
                (8, 3): subtotal_presentation,
            },
        )

        workbook = load_workbook(BytesIO(build_audit_workbook([table])), data_only=False)
        values = [cell.value for row in workbook["T001_TM"] for cell in row]

        self.assertIn("=ROUND(B9-SUM(B7:B8),2)", values)
        self.assertIn("=ROUND(B12-SUM(B11:B11),2)", values)
        self.assertIn("=ROUND(B14-B9+B12,2)", values)
        self.assertIn("=ROUND(D14-D9+D12,2)", values)

    def test_deduction_sections_check_both_subtotals_and_final_total(self):
        subtotal_presentation = CellPresentation(
            bold=False,
            bottom_border=CellBorder(style="double", color="1F1F1F"),
        )
        table = ExtractedTable(
            1,
            [
                ["", "", "Năm nay", "", "Năm trước"],
                ["", "", "VND", "", "VND"],
                ["", "", "", "", ""],
                ["Khoản vay A", "", "100", "", "400"],
                ["Khoản vay B", "", "80", "", "200"],
                ["Khoản vay C", "", "60", "", "100"],
                ["Khoản vay D", "", "35", "", "9"],
                ["", "", "275", "", "709"],
                ["", "", "", "", ""],
                ["Trừ:", "", "", "", ""],
                ["Nợ phải trả trong vòng 12 tháng, trong đó:", "", "", "", ""],
                ["Ngân hàng A", "", "-", "", "35"],
                ["Ngân hàng B", "", "-", "", "36"],
                ["Ngân hàng C", "", "90", "", "-"],
                ["", "", "90", "", "71"],
                ["", "", "", "", ""],
                ["Tổng cộng", "", "185", "", "638"],
            ],
            StatementType.NOTE,
            "Dài hạn",
            cell_presentations={
                (7, 2): subtotal_presentation,
                (7, 4): subtotal_presentation,
                (14, 2): subtotal_presentation,
                (14, 4): subtotal_presentation,
            },
        )

        workbook = load_workbook(BytesIO(build_audit_workbook([table])), data_only=False)
        values = [cell.value for row in workbook["T001_TM"] for cell in row]

        self.assertIn("=ROUND(C11-SUM(C7:C10),2)", values)
        self.assertIn("=ROUND(E11-SUM(E7:E10),2)", values)
        self.assertIn("=ROUND(C18-SUM(C15:C17),2)", values)
        self.assertIn("=ROUND(E18-SUM(E15:E17),2)", values)
        self.assertIn("=ROUND(C20-C11+C18,2)", values)
        self.assertIn("=ROUND(E20-E11+E18,2)", values)

    def test_fixed_asset_net_values_reconcile_cost_and_accumulated_depreciation(self):
        table = ExtractedTable(
            1,
            [
                ["", "Nhà cửa", "Máy móc", "Tổng cộng"],
                ["", "VND", "VND", "VND"],
                ["", "", "", ""],
                ["Nguyên giá", "", "", ""],
                ["Số đầu năm", "595.742.347.364", "445.219.518.167", "1.052.211.452.586"],
                ["Tăng trong năm", "138.464.844.998", "223.502.029.573", "363.111.565.249"],
                ["Phân loại lại", "-", "2.459.960.000", "-"],
                ["Số cuối năm", "734.207.192.362", "671.181.507.740", "1.415.323.017.835"],
                ["", "", "", ""],
                ["Giá trị hao mòn", "", "", ""],
                ["Số đầu năm", "(11.819.931.010)", "(26.002.497.113)", "(38.573.307.915)"],
                ["Tăng trong năm", "(20.348.846.502)", "(44.500.219.187)", "(65.988.736.686)"],
                ["Phân loại lại", "-", "(221.396.400)", "-"],
                ["Số cuối năm", "(32.168.777.512)", "(70.724.112.700)", "(104.562.044.601)"],
                ["", "", "", ""],
                ["Giá trị còn lại", "", "", ""],
                ["Số đầu năm", "583.922.416.354", "419.217.021.054", "1.013.638.144.671"],
                ["", "", "", ""],
                ["Số cuối năm", "702.038.414.850", "600.457.395.040", "1.310.760.973.234"],
            ],
            StatementType.NOTE,
            "Tài sản cố định hữu hình",
        )

        workbook = load_workbook(BytesIO(build_audit_workbook([table])), data_only=False)
        worksheet = workbook["T001_TM"]
        values = [cell.value for row in worksheet for cell in row]

        self.assertIn("=ROUND(B20-B8-B14,2)", values)
        self.assertIn("=ROUND(B22-B11-B17,2)", values)
        self.assertIn("=ROUND(D20-D8-D14,2)", values)
        self.assertIn("=ROUND(D22-D11-D17,2)", values)
        self.assertEqual(worksheet["B14"].number_format, "#,##0;(#,##0);-")
        self.assertEqual(worksheet["E8"].number_format, "#,##0;(#,##0);-")
        self.assertNotIn("[Red]", worksheet["B14"].number_format)
        self.assertNotIn(".##", worksheet["E8"].number_format)
        self.assertGreaterEqual(worksheet.column_dimensions["D"].width, 19)
        self.assertGreaterEqual(worksheet.column_dimensions["E"].width, 19)

    def test_compact_intangible_asset_checks_net_value_and_treats_dash_as_zero(self):
        table = ExtractedTable(
            1,
            [
                ["Phần mềm CDMS", "Số cuối năm", "", "Số đầu năm"],
                ["", "VND", "", "VND"],
                ["", "", "", ""],
                ["Nguyên giá", "350.000.000", "", "-"],
                ["Khấu hao lũy kế", "(69.999.996)", "", "-"],
                ["Giá trị còn lại", "280.000.004", "", "-"],
                ["", "", "", ""],
                ["Khấu hao trong năm", "(69.999.996)", "", "-"],
            ],
            StatementType.NOTE,
            "Tài sản cố định vô hình",
        )

        workbook = load_workbook(BytesIO(build_audit_workbook([table])), data_only=False)
        worksheet = workbook["T001_TM"]
        values = [cell.value for row in worksheet for cell in row]

        self.assertIn("=ROUND(B9-B7-B8,2)", values)
        self.assertIn("=ROUND(D9-D7-D8,2)", values)
        self.assertEqual(worksheet["D7"].value, 0)
        self.assertEqual(worksheet["D8"].value, 0)
        self.assertEqual(worksheet["D9"].value, 0)

    def test_fixed_asset_net_values_accept_depreciation_section_label(self):
        table = ExtractedTable(
            1,
            [
                ["", "Phần mềm", "Tổng cộng"],
                ["", "VND", "VND"],
                ["Nguyên giá", "", ""],
                ["Số đầu năm", "733.064.179", "733.064.179"],
                ["Tăng trong năm", "-", "2.759.348.880"],
                ["Số cuối năm", "733.064.179", "3.492.413.059"],
                ["", "", ""],
                ["Giá trị khấu hao", "", ""],
                ["Số đầu năm", "(75.483.323)", "(75.483.323)"],
                ["Tăng trong năm", "(146.810.760)", "(514.723.944)"],
                ["Số cuối năm", "(222.294.083)", "(590.207.267)"],
                ["", "", ""],
                ["Giá trị còn lại", "", ""],
                ["Số đầu năm", "657.580.856", "657.580.856"],
                ["", "", ""],
                ["Số cuối năm", "510.770.096", "2.902.205.792"],
            ],
            StatementType.NOTE,
            "Tài sản cố định vô hình",
        )

        workbook = load_workbook(BytesIO(build_audit_workbook([table])), data_only=False)
        values = [cell.value for row in workbook["T001_TM"] for cell in row]

        self.assertIn("=ROUND(B17-B7-B12,2)", values)
        self.assertIn("=ROUND(C17-C7-C12,2)", values)
        self.assertIn("=ROUND(B19-B9-B14,2)", values)
        self.assertIn("=ROUND(C19-C9-C14,2)", values)
        self.assertNotIn("=B19-SUM(B17:B17)", values)

    def test_fixed_asset_note_reconciliation_links_all_three_blocks_and_two_periods(self):
        statement = ExtractedTable(
            1,
            [
                ["Chỉ tiêu", "Mã số", "Thuyết minh", "Năm nay", "Năm trước"],
                ["TSCĐ hữu hình", "221", "V.1", "80", "70"],
                ["Nguyên giá", "222", "", "100", "90"],
                ["Hao mòn lũy kế", "223", "", "(20)", "(20)"],
            ],
            StatementType.BALANCE_SHEET_ASSETS,
        )

        def total_row(label, value):
            return [label, "", "", "", "", "", "", "", "", "", "", value]

        note = ExtractedTable(
            2,
            [
                ["", "", "", "", "", "", "", "", "", "", "", "Tổng cộng"],
                ["Nguyên giá"],
                [""],
                [""],
                total_row("Số đầu năm", "90"),
                [""],
                [""],
                total_row("Số cuối năm", "100"),
                ["Giá trị hao mòn"],
                [""],
                total_row("Số đầu năm", "(20)"),
                [""],
                [""],
                total_row("Số cuối năm", "(20)"),
                ["Giá trị còn lại"],
                [""],
                total_row("Số đầu năm", "70"),
                [""],
                total_row("Số cuối năm", "80"),
            ],
            StatementType.NOTE,
            "Tài sản cố định hữu hình",
        )

        workbook = load_workbook(BytesIO(build_audit_workbook([statement, note])), data_only=False)
        self.assertNotIn("01_Doi_chieu_TM", workbook.sheetnames)
        self.assertEqual(workbook["00_Tong_hop"]["J3"].value, "Khớp")

    def test_narrative_loan_contract_table_is_marked_not_checked(self):
        table = ExtractedTable(
            1,
            [
                ["Số hợp đồng", "250828-TFB01620598"],
                ["Ngày hợp đồng", "28/08/2025"],
                ["Thời hạn", "1 năm"],
                ["Mục đích", "Hỗ trợ dịch vụ thanh toán và tài trợ thương mại"],
                ["Lãi suất", "Theo từng lần giải ngân"],
                ["Tài sản thế chấp", "Không"],
                ["Số dư cuối năm", "830.300 USD (tương đương 21.784.581.100 VND)"],
            ],
            StatementType.NOTE,
            "Chi tiết khoản vay",
        )

        workbook = load_workbook(BytesIO(build_audit_workbook([table])), data_only=False)

        self.assertEqual(workbook["00_Tong_hop"]["I2"].value, "Không kiểm tra")
        self.assertIn("không có cấu trúc tổng", workbook["00_Tong_hop"]["K2"].value)
        self.assertEqual(workbook["T001_TM"]["A2"].value, "Tiếp →")

    def test_related_party_relationship_table_is_marked_not_checked(self):
        table = ExtractedTable(
            1,
            [
                ["Bên liên quan", "Mối quan hệ"],
                ["", ""],
                ["Licona Corporation", "Chủ sở hữu"],
            ],
            StatementType.NOTE,
            "Giao dịch với các bên liên quan",
        )

        workbook_bytes, comment_targets = build_audit_workbook_with_comment_targets([table])
        workbook = load_workbook(BytesIO(workbook_bytes), data_only=False)

        self.assertEqual(workbook["00_Tong_hop"]["C1"].value, "Nội dung bảng")
        self.assertEqual(workbook["00_Tong_hop"]["D1"].value, "Loại bảng")
        self.assertEqual(workbook["00_Tong_hop"]["C2"].value, "Giao dịch với các bên liên quan")
        self.assertEqual(
            workbook["00_Tong_hop"]["D2"].value,
            "Giao dịch với các bên liên quan",
        )
        self.assertEqual(workbook["00_Tong_hop"]["I2"].value, "Không kiểm tra")
        self.assertFalse(comment_targets)
        self.assertIn(
            "Không áp dụng đối chiếu BCTC cho bảng này.",
            [cell.value for row in workbook["T001_TM"].iter_rows() for cell in row],
        )

    def test_tax_rollforward_reconciliation_is_marked_not_applicable(self):
        table = ExtractedTable(
            1,
            [
                ["", "Số đầu năm", "Số phải nộp", "Số đã nộp, khấu trừ", "Số cuối năm"],
                ["", "VND", "VND", "VND", "VND"],
                ["Thuế giá trị gia tăng", "-", "100", "(100)", "-"],
                ["Thuế thu nhập cá nhân", "20", "30", "(10)", "40"],
                ["Tổng cộng", "20", "130", "(110)", "40"],
            ],
            StatementType.NOTE,
            "Thuế và các khoản phải nộp Nhà nước",
        )

        workbook_bytes, comment_targets = build_audit_workbook_with_comment_targets([table])
        workbook = load_workbook(BytesIO(workbook_bytes), data_only=False)

        self.assertEqual(workbook["00_Tong_hop"]["J2"].value, "Không áp dụng")
        self.assertFalse(comment_targets)
        self.assertIn(
            "Không áp dụng đối chiếu BCTC cho bảng này.",
            [cell.value for row in workbook["T001_TM"].iter_rows() for cell in row],
        )

    def test_related_party_amount_tables_without_totals_are_not_checked_or_commented(self):
        tables = [
            ExtractedTable(
                1,
                [
                    ["", "Năm nay", "", "Năm trước"],
                    ["", "VND", "", "VND"],
                    ["Mượn tiền", "", "", ""],
                    ["Ông A", "2.500.000.000", "", "26.870.000.000"],
                    ["Trả tiền", "", "", ""],
                    ["Ông A", "20.000.000.000", "", "20.103.424.847"],
                ],
                StatementType.NOTE,
                "Các nghiệp vụ kinh tế quan trọng với các bên liên quan",
                context_hints=("Giao dịch với các bên liên quan",),
            ),
            ExtractedTable(
                2,
                [
                    ["", "Số cuối năm", "", "Số đầu năm"],
                    ["", "VND", "", "VND"],
                    ["Trả trước cho người bán", "", "", ""],
                    ["Công ty A", "-", "", "750.000.000"],
                    ["Phải trả khác", "", "", ""],
                    ["Ông A", "21.605.556.052", "", "38.805.556.052"],
                ],
                StatementType.NOTE,
                "Các nghiệp vụ kinh tế quan trọng với các bên liên quan",
                context_hints=("Giao dịch với các bên liên quan",),
            ),
        ]

        workbook_bytes, comment_targets = build_audit_workbook_with_comment_targets(tables)
        workbook = load_workbook(BytesIO(workbook_bytes), data_only=False)

        for row in (2, 3):
            self.assertEqual(workbook["00_Tong_hop"].cell(row, 3).value, "Bên liên quan")
            self.assertEqual(workbook["00_Tong_hop"].cell(row, 4).value, "Thuyết minh")
            self.assertEqual(workbook["00_Tong_hop"].cell(row, 9).value, "Không kiểm tra")
        self.assertFalse(comment_targets)
        for sheet_name in ("T001_TM", "T002_TM"):
            self.assertIn(
                "Không áp dụng đối chiếu BCTC cho bảng này.",
                [cell.value for row in workbook[sheet_name].iter_rows() for cell in row],
            )

    def test_unrecognized_note_table_still_requires_review(self):
        table = ExtractedTable(
            1,
            [["Chỉ tiêu chưa rõ", "100"], ["Dòng khác", "200"]],
            StatementType.NOTE,
            "Bảng chưa xác định",
        )

        workbook = load_workbook(BytesIO(build_audit_workbook([table])), data_only=False)

        self.assertEqual(workbook["00_Tong_hop"]["I2"].value, "Cần xem xét")

    def test_unknown_qualitative_table_is_not_checked_or_commented(self):
        table = ExtractedTable(
            1,
            [["Nội dung mô tả", "Không áp dụng"], ["Thông tin khác", "Có"]],
            StatementType.UNKNOWN,
            "Thông tin chưa có rule riêng",
        )

        workbook_bytes, comment_targets = build_audit_workbook_with_comment_targets([table])
        workbook = load_workbook(BytesIO(workbook_bytes), data_only=False)

        self.assertEqual(workbook["00_Tong_hop"]["I2"].value, "Không kiểm tra")
        self.assertFalse(comment_targets)

    def test_unknown_table_with_accounting_values_still_requires_review(self):
        table = ExtractedTable(
            1,
            [["Chỉ tiêu chưa rõ", "100"], ["Dòng khác", "200"]],
            StatementType.UNKNOWN,
            "Bảng chưa xác định",
        )

        workbook_bytes, comment_targets = build_audit_workbook_with_comment_targets([table])
        workbook = load_workbook(BytesIO(workbook_bytes), data_only=False)

        self.assertEqual(workbook["00_Tong_hop"]["I2"].value, "Cần xem xét")
        self.assertFalse(comment_targets)

    def test_noise_tables_from_word_do_not_create_manual_review_comments(self):
        tables = [
            ExtractedTable(
                1,
                [
                    ["", "", "Tỉ lệ vốn góp (%)"],
                    ["Kingfa Sci. And Tech. Co., Ltd.", "Thành lập tại Trung Quốc", "70"],
                    ["Hongkong Kingfa Development Co., Limited", "Thành lập tại Trung Quốc", "30"],
                ],
                StatementType.NOTE,
                "Các chủ đầu tư của Công ty gồm",
            ),
            ExtractedTable(
                2,
                [
                    ["", "Số cuối năm", "", "Số đầu năm"],
                    ["", "VND", "", "VND"],
                    ["Thuế và các khoản khác phải thu Nhà nước", "460.884.903", "", "423.652.174"],
                    ["Thuế và các khoản phải nộp Nhà nước", "158.255.605", "", "120.392.168"],
                ],
                StatementType.NOTE,
                "Thuế và các khoản phải nộp Nhà nước",
            ),
        ]

        workbook_bytes, comment_targets = build_audit_workbook_with_comment_targets(tables)
        workbook = load_workbook(BytesIO(workbook_bytes), data_only=False)

        self.assertEqual(workbook["00_Tong_hop"]["I2"].value, "Không kiểm tra")
        self.assertEqual(workbook["00_Tong_hop"]["I3"].value, "Không kiểm tra")
        self.assertFalse(
            any(
                target.message == "Kiểm tra số học: Cần xem xét thủ công."
                for target in comment_targets
            )
        )

    def test_informational_note_table_patterns_are_marked_not_checked(self):
        tables = (
            ExtractedTable(
                1,
                [
                    ["", "Số năm khấu hao"],
                    ["Nhà cửa, vật kiến trúc", "11 - 35"],
                    ["Máy móc, thiết bị", "9 - 11"],
                    ["Phương tiện vận tải", "6 - 11"],
                ],
                StatementType.NOTE,
                "Thời gian khấu hao",
            ),
            ExtractedTable(
                2,
                [
                    ["", "VND"],
                    [
                        "Giá trị còn lại cuối năm của tài sản cố định hữu hình đã dùng để thế chấp, cầm cố đảm bảo các khoản vay",
                        "222.547.876.933",
                    ],
                ],
                StatementType.NOTE,
                "Tài sản thế chấp",
            ),
            ExtractedTable(
                3,
                [
                    ["", "Số cuối năm", "Số đầu năm"],
                    ["Ngoại tệ các loại (USD)", "891.035,22", "808.493,30"],
                ],
                StatementType.NOTE,
                "Ngoại tệ",
            ),
        )

        for table in tables:
            with self.subTest(table=table.index):
                workbook_bytes, comment_targets = build_audit_workbook_with_comment_targets([table])
                workbook = load_workbook(BytesIO(workbook_bytes), data_only=False)
                self.assertEqual(workbook["00_Tong_hop"]["I2"].value, "Không kiểm tra")
                if table.index == 3:
                    self.assertEqual(
                        workbook["00_Tong_hop"]["J2"].value,
                        "Không áp dụng",
                    )
                    self.assertIn(
                        "Không áp dụng đối chiếu BCTC cho bảng này.",
                        [
                            cell.value
                            for row in workbook["T003_TM"].iter_rows()
                            for cell in row
                        ],
                    )
                self.assertFalse(
                    any(
                        target.message == "Kiểm tra số học: Cần xem xét thủ công."
                        for target in comment_targets
                    )
                )

    def test_noise_disclosures_do_not_create_arithmetic_comments(self):
        tables = (
            ExtractedTable(
                1,
                [
                    ["", "Năm"],
                    ["Nhà cửa, vật kiến trúc", "5 - 7"],
                    ["Máy móc, thiết bị", "5 - 10"],
                    ["Phương tiện vận tải, truyền dẫn", "5"],
                    ["Thiết bị văn phòng", "3 - 7"],
                ],
                StatementType.UNKNOWN,
                "Chính sách khấu hao tài sản cố định",
            ),
            ExtractedTable(
                2,
                [
                    ["", "VND"],
                    [
                        "Nguyên giá tài sản cố định hữu hình cuối năm đã khấu hao hết nhưng vẫn còn sử dụng",
                        "26.805.684.650",
                    ],
                ],
                StatementType.NOTE,
                "Tài sản cố định hữu hình",
            ),
            ExtractedTable(
                3,
                [
                    ["", "Ông Noh Byung Hwi", ""],
                    ["Số hợp đồng", "20221111/HĐVTCN", "202212-001"],
                    ["Ngày hợp đồng", "11/11/2022", "10/12/2022"],
                    ["Ngày đáo hạn", "11/11/2025", "10/12/2025"],
                    ["Hạn mức cho vay (VND)", "5.000.000.000", "4.030.000.000"],
                    ["Tài sản đảm bảo", "Không", "Không"],
                    ["Số dư cuối kỳ (VND)", "2.981.820.000", "4.030.000.000"],
                ],
                StatementType.NOTE,
                "Vay và nợ thuê tài chính dài hạn",
            ),
        )

        workbook_bytes, comment_targets = build_audit_workbook_with_comment_targets(list(tables))
        workbook = load_workbook(BytesIO(workbook_bytes), data_only=False)

        self.assertTrue(
            all(workbook["00_Tong_hop"].cell(row, 9).value == "Không kiểm tra" for row in range(2, 5))
        )
        self.assertFalse(comment_targets)
        self.assertFalse(
            any(
                isinstance(cell.value, str) and cell.value.startswith("=ROUND(")
                for sheet_name in ("T001_KHAC", "T002_TM", "T003_TM")
                for row in workbook[sheet_name].iter_rows()
                for cell in row
            )
        )

    def test_explicit_total_row_with_one_component_receives_vertical_checks(self):
        table = ExtractedTable(
            1,
            [
                ["", "Vốn đầu tư", "", "Số cuối năm", "", "Vốn đã góp", "", "Số đầu năm"],
                ["", "USD", "", "USD", "", "VND", "", "VND"],
                ["", "", "", "", "", "", "", ""],
                ["Thrive Nation Group Limited", "1.000", "", "35.500.000", "", "900", "", "529.677.417.162"],
                ["Tổng cộng", "1.000", "", "36.500.000", "", "900", "", "529.677.417.162"],
            ],
            StatementType.NOTE,
            "Theo giấy chứng nhận đăng ký đầu tư",
        )

        workbook_bytes, comment_targets = build_audit_workbook_with_comment_targets([table])
        workbook = load_workbook(BytesIO(workbook_bytes), data_only=False)
        values = [cell.value for row in workbook["T001_TM"] for cell in row]

        self.assertIn("Kiểm tra cộng dọc dòng 8", values)
        self.assertIn("=ROUND(B8-SUM(B7:B7),2)", values)
        self.assertIn("=ROUND(D8-SUM(D7:D7),2)", values)
        self.assertIn("=ROUND(F8-SUM(F7:F7),2)", values)
        self.assertIn("=ROUND(H8-SUM(H7:H7),2)", values)
        self.assertIn(
            "Kiểm tra số học: chênh lệch 1.000.000.",
            {target.message for target in comment_targets},
        )

    def test_note_status_labels_are_vietnamese(self):
        self.assertEqual(
            NOTE_STATUS_LABELS,
            {
                "Matched": "Khớp",
                "Difference": "Có chênh lệch",
                "Note value not found": "Không tìm thấy số liệu TM",
                "Statement not found": "Không tìm thấy chỉ tiêu BCTC",
            },
        )

    def test_summary_sort_options_and_detail_sheet_order_are_independent(self):
        tables = [
            ExtractedTable(3, [["Bảng 3"]], StatementType.UNKNOWN),
            ExtractedTable(1, [["Bảng 1"]], StatementType.UNKNOWN),
            ExtractedTable(2, [["Bảng 2"]], StatementType.UNKNOWN),
        ]
        results = [
            TableCheckResult(1, "Không kiểm tra", issue_count=0),
            TableCheckResult(2, "Cần xem xét", issue_count=1),
            TableCheckResult(3, "Có kiểm tra", issue_count=3),
        ]

        by_issues = _sort_summary_tables(tables, results, SUMMARY_SORT_ISSUES_DESC)
        by_status = _sort_summary_tables(tables, results, SUMMARY_SORT_STATUS)
        workbook = load_workbook(
            BytesIO(build_audit_workbook(tables, summary_sort=SUMMARY_SORT_STATUS))
        )
        default_workbook = load_workbook(BytesIO(build_audit_workbook(tables)))

        self.assertEqual([table.index for table in by_issues], [3, 2, 1])
        self.assertEqual([table.index for table in by_status], [2, 3, 1])
        self.assertEqual(workbook.sheetnames[1:], ["T001_KHAC", "T002_KHAC", "T003_KHAC"])
        self.assertEqual(
            [default_workbook["00_Tong_hop"].cell(row, 1).value for row in range(2, 5)],
            [1, 2, 3],
        )

    def test_detail_backlinks_follow_sorted_summary_rows(self):
        tables = [
            ExtractedTable(1, [["Thông tin"]], StatementType.GENERAL_INFO),
            ExtractedTable(2, [["Chữ ký"]], StatementType.SIGNATURE),
            ExtractedTable(3, [["Chưa phân loại"]], StatementType.UNKNOWN),
        ]

        workbook = load_workbook(
            BytesIO(build_audit_workbook(tables, summary_sort=SUMMARY_SORT_STATUS))
        )

        self.assertEqual(
            [workbook["00_Tong_hop"].cell(row, 1).value for row in range(2, 5)],
            [1, 2, 3],
        )
        self.assertEqual(workbook["T001_TT"]["B2"].hyperlink.target, "#00_Tong_hop!B2")
        self.assertEqual(workbook["T002_CK"]["B2"].hyperlink.target, "#00_Tong_hop!B3")
        self.assertEqual(workbook["T003_KHAC"]["B2"].hyperlink.target, "#00_Tong_hop!B4")

    def test_detail_sheets_have_sequential_navigation_in_source_order(self):
        tables = [
            ExtractedTable(3, [["Bảng 3"]], StatementType.UNKNOWN),
            ExtractedTable(1, [["Bảng 1"]], StatementType.UNKNOWN),
            ExtractedTable(2, [["Bảng 2"]], StatementType.UNKNOWN),
        ]

        workbook = load_workbook(BytesIO(build_audit_workbook(tables)))

        first = workbook["T001_KHAC"]
        middle = workbook["T002_KHAC"]
        last = workbook["T003_KHAC"]
        self.assertEqual(first["B1"].value, "Chưa phân loại")
        self.assertEqual(first["B2"].value, "Trở về Tổng hợp")
        self.assertIsNone(first["A1"].hyperlink)
        self.assertEqual(first["A2"].hyperlink.target, "#T002_KHAC!A1")
        self.assertEqual(middle["A1"].hyperlink.target, "#T001_KHAC!A1")
        self.assertEqual(middle["A2"].hyperlink.target, "#T003_KHAC!A1")
        self.assertEqual(last["A1"].hyperlink.target, "#T002_KHAC!A1")
        self.assertIsNone(last["A2"].hyperlink)

    def test_check_cells_and_review_legend_use_distinct_visual_roles(self):
        table = ExtractedTable(
            1,
            [
                ["Chỉ tiêu", "Mã", "Năm nay", "Năm trước"],
                ["Doanh thu bán hàng", "01", "100", "90"],
                ["Các khoản giảm trừ", "02", "10", "5"],
                ["Doanh thu thuần", "10", "90", "85"],
            ],
            StatementType.INCOME_STATEMENT,
        )

        workbook = load_workbook(BytesIO(build_audit_workbook([table])), data_only=False)
        worksheet = workbook["T001_PL"]
        formula_cells = [
            cell
            for row in worksheet.iter_rows()
            for cell in row
            if isinstance(cell.value, str) and cell.value.startswith("=ROUND(")
        ]

        self.assertTrue(formula_cells)
        self.assertTrue(
            all(str(cell.fill.fgColor.rgb).endswith("F2F2F2") for cell in formula_cells)
        )
        self.assertIsNone(worksheet["A8"].fill.fill_type)
        self.assertEqual(worksheet["A2"].value, "Tiếp →")
        self.assertNotIn(
            "Tổng sai lệch tuyệt đối",
            [cell.value for row in worksheet.iter_rows() for cell in row],
        )
        self.assertNotIn(
            "Tổng sai sót",
            [cell.value for row in worksheet.iter_rows() for cell in row],
        )
        self.assertNotIn(
            "Tổng sai lệch tuyệt đối",
            [cell.value for row in workbook["00_Tong_hop"].iter_rows() for cell in row],
        )

    def test_invalid_sort_option_is_rejected(self):
        with self.assertRaisesRegex(ValueError, "00_Tong_hop"):
            build_audit_workbook([], summary_sort="khong_hop_le")

    def test_shared_color_roles_are_applied_to_all_sheet_types(self):
        tables = [
            ExtractedTable(
                index=index,
                rows=[
                    ["CHỈ TIÊU", "Mã số", "Thuyết minh", "Năm nay", "Năm trước"],
                    ["Chỉ tiêu thử nghiệm", str(index), "", "100", "90"],
                ],
                statement_type=statement_type,
                title_hint=statement_type.value,
            )
            for index, statement_type in enumerate(
                (
                    StatementType.BALANCE_SHEET_ASSETS,
                    StatementType.INCOME_STATEMENT,
                    StatementType.CASH_FLOW,
                    StatementType.NOTE,
                ),
                start=1,
            )
        ]

        workbook = load_workbook(BytesIO(build_audit_workbook(tables)))

        self.assertTrue(workbook["00_Tong_hop"]["A1"].fill.fgColor.rgb.endswith(COLOR_NAVY))
        self.assertNotIn("01_Doi_chieu_TM", workbook.sheetnames)
        for sheet_name in ("T001_BS_TS", "T002_PL", "T003_CF", "T004_TM"):
            worksheet = workbook[sheet_name]
            self.assertTrue(worksheet["B1"].fill.fgColor.rgb.endswith(COLOR_NAVY))
            self.assertTrue(worksheet["A3"].fill.fgColor.rgb.endswith(COLOR_BLUE))

    def test_same_table_arithmetic_marks_invalid_period_for_review(self):
        table = ExtractedTable(
            1,
            [
                ["Chỉ tiêu", "Mã số", "Năm nay", "Năm trước"],
                ["Doanh thu", "1", "100", "90"],
                ["Khoản giảm trừ", "2", "không hợp lệ", "10"],
                ["Doanh thu thuần", "10", "100", "80"],
            ],
            StatementType.INCOME_STATEMENT,
        )

        workbook = load_workbook(BytesIO(build_audit_workbook([table])), data_only=False)
        worksheet = workbook["T001_PL"]

        self.assertEqual(worksheet["E7"].value, 0)
        self.assertEqual(worksheet["F7"].value, "=ROUND(D7-(D5-D6),2)")
        self.assertIn("Cần xác minh dữ liệu", worksheet["G7"].value)
        self.assertIn("kỳ này: mã 2", worksheet["G7"].value)
        self.assertTrue(worksheet["C7"].fill.fgColor.rgb.endswith("FFF2CC"))
        self.assertEqual(workbook["00_Tong_hop"]["I2"].value, "Cần xem xét")
        self.assertIn("dữ liệu thiếu hoặc không hợp lệ", workbook["00_Tong_hop"]["K2"].value)
        self.assertTrue(workbook["00_Tong_hop"]["I2"].fill.fgColor.rgb.endswith("FFF2CC"))

    def test_missing_target_comments_only_actionable_component_values(self):
        table = ExtractedTable(
            1,
            [
                ["Chỉ tiêu", "Mã số", "Năm nay", "Năm trước"],
                ["Doanh thu", "1", "100", "0"],
                ["Khoản giảm trừ", "2", "-", "không hợp lệ"],
            ],
            StatementType.INCOME_STATEMENT,
        )

        workbook = load_workbook(BytesIO(build_audit_workbook([table])), data_only=False)
        worksheet = workbook["T001_PL"]

        self.assertIsNotNone(worksheet["C5"].comment)
        self.assertIn("thiếu mã đích 10", worksheet["C5"].comment.text)
        self.assertIsNone(worksheet["D5"].comment)
        self.assertIsNone(worksheet["C6"].comment)
        self.assertIsNotNone(worksheet["D6"].comment)
        self.assertEqual(workbook["00_Tong_hop"]["I2"].value, "Cần xem xét")

        _, comment_targets = build_audit_workbook_with_comment_targets([table])
        target_locations = {
            (target.row_index, target.column_index): target.message
            for target in comment_targets
        }
        self.assertEqual(set(target_locations), {(1, 2), (2, 3)})
        self.assertTrue(
            all("thiếu mã đích 10" in message for message in target_locations.values())
        )

    def test_note_sheet_contains_its_reconciliation_results(self):
        statement = ExtractedTable(
            index=1,
            rows=[
                ["CHỈ TIÊU", "Mã số", "Thuyết minh", "Năm nay", "Năm trước"],
                ["Hàng tồn kho", "140", "V.6", "1.000", "900"],
            ],
            statement_type=StatementType.BALANCE_SHEET_ASSETS,
        )
        note = ExtractedTable(
            index=2,
            rows=[
                ["", "Năm nay", "Năm trước"],
                ["Tổng cộng", "1.000", "900"],
            ],
            statement_type=StatementType.NOTE,
            title_hint="Hàng tồn kho",
        )

        workbook = load_workbook(BytesIO(build_audit_workbook([statement, note])), data_only=False)
        worksheet = workbook["T002_TM"]
        values = [cell.value for row in worksheet.iter_rows() for cell in row]

        self.assertIn("ĐỐI CHIẾU VỚI BCTC", values)
        self.assertNotIn("140", values)
        self.assertNotIn("Khớp", values)
        self.assertNotIn("Trạng thái", values)
        self.assertNotIn("Matched", values)
        bctc_row = next(cell.row for cell in worksheet["A"] if cell.value == "Giá trị BCTC")
        self.assertEqual(worksheet.cell(bctc_row, 2).value, "='T001_BS_TS'!D5")
        self.assertEqual(worksheet.cell(bctc_row, 3).value, "='T001_BS_TS'!E5")
        self.assertEqual(
            worksheet.cell(bctc_row + 1, 2).value,
            f"=ROUND(N(B{bctc_row})-N(B5),2)",
        )
        self.assertEqual(
            worksheet.cell(bctc_row + 1, 3).value,
            f"=ROUND(N(C{bctc_row})-N(C5),2)",
        )
        self.assertEqual(worksheet.cell(bctc_row, 2).hyperlink.target, "#T001_BS_TS!D5")

        self.assertNotIn("01_Doi_chieu_TM", workbook.sheetnames)
        self.assertEqual(workbook["00_Tong_hop"]["J3"].value, "Khớp")
        self.assertFalse(
            any(
                cell.hyperlink and cell.hyperlink.target == "#01_Doi_chieu_TM!A3"
                for row in worksheet
                for cell in row
            )
        )

    def test_date_range_periods_write_distinct_reconciliation_columns(self):
        statement = ExtractedTable(
            1,
            [
                ["CHỈ TIÊU", "Mã số", "Thuyết minh", "Năm nay", "Năm trước"],
                ["14. Tổng lợi nhuận kế toán trước thuế", "50", "", "(6.139.421.945)", "(3.305.515.469)"],
            ],
            StatementType.INCOME_STATEMENT,
        )
        note = ExtractedTable(
            2,
            [
                ["", "Từ 01/01/2025 đến 31/12/2025", "", "Từ 14/05/2024 đến 31/12/2024"],
                ["", "VND", "", "VND"],
                ["Lợi nhuận trước thuế", "(6.139.421.945)", "", "(3.305.515.469)"],
            ],
            StatementType.NOTE,
        )

        workbook = load_workbook(BytesIO(build_audit_workbook([statement, note])), data_only=False)
        worksheet = workbook["T002_TM"]
        bctc_row = next(cell.row for cell in worksheet["A"] if cell.value == "Giá trị BCTC")

        self.assertEqual(worksheet.cell(bctc_row, 2).value, "='T001_PL'!D5")
        self.assertEqual(worksheet.cell(bctc_row, 4).value, "='T001_PL'!E5")
        self.assertEqual(
            worksheet.cell(bctc_row + 1, 2).value,
            f"=ROUND(N(B{bctc_row})-N(B6),2)",
        )
        self.assertEqual(
            worksheet.cell(bctc_row + 1, 4).value,
            f"=ROUND(N(D{bctc_row})-N(D6),2)",
        )

    def test_tt200_equity_movement_writes_three_traceable_reconciliation_blocks(self):
        balance_sheet = ExtractedTable(
            1,
            [
                ["CHỈ TIÊU", "Mã số", "Thuyết minh", "Năm nay", "Năm trước"],
                ["Vốn góp của chủ sở hữu", "411", "V.16", "100", "90"],
                ["Lợi nhuận sau thuế chưa phân phối", "421", "V.16", "50", "40"],
            ],
            StatementType.BALANCE_SHEET_EQUITY,
        )
        income_statement = ExtractedTable(
            2,
            [
                ["CHỈ TIÊU", "Mã số", "Thuyết minh", "Năm nay", "Năm trước"],
                ["Lợi nhuận sau thuế thu nhập doanh nghiệp", "60", "", "10", "8"],
            ],
            StatementType.INCOME_STATEMENT,
        )
        note = ExtractedTable(
            3,
            [
                ["", "Vốn góp của chủ sở hữu", "", "Lợi nhuận sau thuế chưa phân phối"],
                ["Số đầu năm trước", "80", "", "32"],
                ["Lợi nhuận thuần trong năm trước", "-", "", "8"],
                ["Số cuối năm trước, đầu năm nay", "90", "", "40"],
                ["Lợi nhuận thuần trong năm", "-", "", "10"],
                ["Số cuối năm", "100", "", "50"],
            ],
            StatementType.NOTE,
            "16. Vốn chủ sở hữu",
        )

        workbook = load_workbook(
            BytesIO(build_audit_workbook([balance_sheet, income_statement, note])),
            data_only=False,
        )
        worksheet = workbook["T003_TM"]
        values = [cell.value for row in worksheet.iter_rows() for cell in row]
        links = {
            cell.hyperlink.target
            for row in worksheet.iter_rows()
            for cell in row
            if cell.hyperlink
        }

        self.assertIn("Đối chiếu 1: Vốn góp của chủ sở hữu", values)
        self.assertIn("Đối chiếu 2: Lợi nhuận sau thuế chưa phân phối", values)
        self.assertIn("Đối chiếu 3: Lợi nhuận sau thuế thu nhập doanh nghiệp", values)
        self.assertIn("#T001_BS_NV!D5", links)
        self.assertIn("#T001_BS_NV!D6", links)
        self.assertIn("#T002_PL!D5", links)
        self.assertTrue(any(value and "-N(B9)" in str(value) for value in values))
        self.assertTrue(any(value and "-N(D8)" in str(value) for value in values))
        self.assertTrue(any(value and "-N(D7)" in str(value) for value in values))

    def test_each_note_sheet_only_contains_its_own_match(self):
        statement = ExtractedTable(
            index=1,
            rows=[
                ["CHỈ TIÊU", "Mã số", "Thuyết minh", "Năm nay", "Năm trước"],
                ["Tiền", "110", "V.1", "100", "90"],
                ["Hàng tồn kho", "140", "V.6", "200", "180"],
            ],
            statement_type=StatementType.BALANCE_SHEET_ASSETS,
        )
        notes = [
            ExtractedTable(2, [["Tổng cộng", "100", "90"]], StatementType.NOTE, "Tiền"),
            ExtractedTable(3, [["Tổng cộng", "200", "180"]], StatementType.NOTE, "Hàng tồn kho"),
        ]

        workbook = load_workbook(BytesIO(build_audit_workbook([statement, *notes])))
        first_values = [cell.value for row in workbook["T002_TM"] for cell in row]
        second_values = [cell.value for row in workbook["T003_TM"] for cell in row]

        self.assertNotIn("110", first_values)
        self.assertNotIn("140", first_values)
        self.assertNotIn("140", second_values)
        self.assertNotIn("110", second_values)
        first_link_targets = {
            cell.hyperlink.target
            for row in workbook["T002_TM"]
            for cell in row
            if cell.hyperlink
        }
        second_link_targets = {
            cell.hyperlink.target
            for row in workbook["T003_TM"]
            for cell in row
            if cell.hyperlink
        }
        self.assertIn("#T001_BS_TS!D5", first_link_targets)
        self.assertIn("#T001_BS_TS!D6", second_link_targets)
        self.assertNotIn("#01_Doi_chieu_TM!A3", first_link_targets)
        self.assertNotIn("#01_Doi_chieu_TM!A4", second_link_targets)

    def test_note_sheet_preserves_source_presentation_in_separate_columns(self):
        note = ExtractedTable(
            index=1,
            rows=[
                ["Tiêu đề", "Tiêu đề", ""],
                ["Tổng cộng", "100", "90"],
            ],
            statement_type=StatementType.NOTE,
            title_hint="Thuyết minh thử nghiệm",
            cell_presentations={
                (0, 0): CellPresentation(
                    bold=True,
                    italic=True,
                    underline=True,
                    fill_color="D9EAF7",
                    horizontal_alignment="center",
                    vertical_alignment="center",
                    bottom_border=CellBorder(style="double", color="1F4E78"),
                )
            },
            merged_ranges=((0, 0, 0, 1),),
            column_widths=(3.0, 1.0, 1.5),
        )

        workbook = load_workbook(BytesIO(build_audit_workbook([note])))
        worksheet = workbook["T001_TM"]
        source_cell = worksheet["A4"]
        reconciliation_title = next(
            cell
            for row in worksheet.iter_rows()
            for cell in row
            if cell.value == "ĐỐI CHIẾU VỚI BCTC"
        )

        self.assertIn("A4:B4", {str(cell_range) for cell_range in worksheet.merged_cells.ranges})
        self.assertTrue(source_cell.font.bold)
        self.assertTrue(source_cell.font.italic)
        self.assertEqual(source_cell.font.underline, "single")
        self.assertEqual(source_cell.fill.fgColor.rgb, "00D9EAF7")
        self.assertEqual(source_cell.alignment.horizontal, "center")
        self.assertEqual(source_cell.border.bottom.style, "double")
        self.assertGreater(worksheet.column_dimensions["A"].width, worksheet.column_dimensions["B"].width)
        self.assertEqual(reconciliation_title.column, 1)
        self.assertIn(
            f"A{reconciliation_title.row}:C{reconciliation_title.row}",
            {str(cell_range) for cell_range in worksheet.merged_cells.ranges},
        )
        self.assertEqual(worksheet.freeze_panes, "B4")

    def test_items_sharing_note_reference_are_written_as_separate_vertical_blocks(self):
        statement = ExtractedTable(
            index=1,
            rows=[
                ["CHỈ TIÊU", "Mã số", "Thuyết minh", "Năm nay", "Năm trước"],
                ["Doanh thu bán hàng", "01", "VI.1", "2.100", "1.600"],
                ["Doanh thu thuần", "10", "VI.1", "2.000", "1.500"],
            ],
            statement_type=StatementType.INCOME_STATEMENT,
        )
        note = ExtractedTable(
            index=2,
            rows=[
                ["", "Năm nay", "Năm trước"],
                ["Tổng cộng", "2.000", "1.500"],
            ],
            statement_type=StatementType.NOTE,
            title_hint="Doanh thu thuần",
        )

        workbook = load_workbook(BytesIO(build_audit_workbook([statement, note])))
        worksheet = workbook["T002_TM"]
        block_headers = [
            cell
            for row in worksheet.iter_rows()
            for cell in row
            if isinstance(cell.value, str) and re.match(r"^Đối chiếu \d+:", cell.value)
        ]
        bctc_labels = [cell for cell in worksheet["A"] if cell.value == "Giá trị BCTC"]
        self.assertEqual(len(block_headers), 2)
        self.assertEqual(len(bctc_labels), 2)
        displayed_values = [cell.value for row in worksheet.iter_rows() for cell in row]
        self.assertNotIn("Khớp", displayed_values)
        self.assertNotIn("Trạng thái", displayed_values)
        self.assertIn("Cần xác minh", displayed_values)
        self.assertIn("Không tìm thấy số liệu TM", displayed_values)
        self.assertNotIn("Thông tin truy vết", displayed_values)
        self.assertNotIn("Cơ sở đối chiếu", displayed_values)
        self.assertNotIn("Xem toàn bộ tại 01_Doi_chieu_TM", displayed_values)
        self.assertNotIn("Nhiều kết quả khớp", displayed_values)
        self.assertNotIn("Multiple matches", displayed_values)

    def test_detail_sheets_are_ordered_by_auditor_priority(self):
        info = ExtractedTable(1, [["Thông tin công ty", "ABC"]], StatementType.GENERAL_INFO)
        signature = ExtractedTable(2, [["Người lập biểu", "Kế toán trưởng"]], StatementType.SIGNATURE)
        bs = ExtractedTable(3, [["Tổng Tài sản", "100", "90"]], StatementType.BALANCE_SHEET_ASSETS)
        pl = ExtractedTable(4, [["Doanh thu thuần", "50", "40"]], StatementType.INCOME_STATEMENT)
        checked_note = ExtractedTable(
            5,
            [["Dòng 1", "10"], ["Dòng 2", "20"], ["Tổng cộng", "30"]],
            StatementType.NOTE,
            title_hint="Thuyết minh số liệu",
        )

        workbook = load_workbook(
            BytesIO(build_audit_workbook([info, signature, bs, pl, checked_note]))
        )

        expected_order = [
            "00_Tong_hop",
            "T003_BS_TS",
            "T004_PL",
            "T005_TM",
            "T001_TT",
            "T002_CK",
        ]
        self.assertEqual(workbook.sheetnames, expected_order)

    def test_sheet_tab_colors_reflect_audit_status(self):
        summary_sheet_table = ExtractedTable(1, [["Thông tin công ty"]], StatementType.GENERAL_INFO)
        error_table = ExtractedTable(
            2,
            [["", "Năm nay"], ["Dòng 1", "10"], ["Dòng 2", "20"], ["Tổng cộng", "999"]],
            StatementType.NOTE,
            title_hint="Chi tiết",
        )
        review_table = ExtractedTable(
            3,
            [["", "Năm nay"], ["Dòng 1", "10"], ["Dòng 2", "20"], ["Tổng cộng", "30"]],
            StatementType.NOTE,
            title_hint="Chi tiết xem xét",
        )

        workbook = load_workbook(
            BytesIO(build_audit_workbook([summary_sheet_table, error_table, review_table]))
        )

        self.assertEqual(workbook["00_Tong_hop"].sheet_properties.tabColor.rgb, "001F4E78")
        self.assertEqual(workbook["T002_TM"].sheet_properties.tabColor.rgb, "00C00000")
        self.assertEqual(workbook["T003_TM"].sheet_properties.tabColor.rgb, "00FFC000")
        self.assertIsNone(workbook["T001_TT"].sheet_properties.tabColor)

    def test_missing_code_verification_comments_produce_yellow_tab_and_review_category(self):
        bs_table = ExtractedTable(
            1,
            [
                ["Chỉ tiêu", "Mã số", "Năm nay", "Năm trước"],
                ["Vốn góp của chủ sở hữu", "411", "100", "100"],
                ["Cổ phiếu phổ thông có quyền biểu quyết", "411a", "100", "100"],
            ],
            StatementType.BALANCE_SHEET_EQUITY,
            title_hint="Bảng cân đối kế toán",
        )
        workbook_bytes, targets = build_audit_workbook_with_comment_targets([bs_table])
        workbook = load_workbook(BytesIO(workbook_bytes), data_only=False)
        review_targets = [t for t in targets if "Cần xác minh" in t.message or "thiếu" in t.message]
        if review_targets:
            for target in review_targets:
                self.assertEqual(target.category, WordCommentCategory.REVIEW)
            # The tab should be YELLOW (00FFC000), NOT RED (00C00000)
            self.assertEqual(workbook["T001_BS_NV"].sheet_properties.tabColor.rgb, "00FFC000")


if __name__ == "__main__":
    unittest.main()


