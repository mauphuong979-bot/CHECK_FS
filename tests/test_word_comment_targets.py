import unittest
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.comments import Comment

from src.domain.models import ExtractedTable, NoteMatchResult, StatementType, TableCheckResult
from src.reporting.excel_report import (
    DATA_START_ROW,
    WARNING_FILL,
    _build_word_comment_targets,
)


class WordCommentTargetTest(unittest.TestCase):
    def test_duplicate_period_reference_does_not_hide_nonzero_difference(self):
        table = ExtractedTable(1, [["Lợi nhuận trước thuế", "100"]], StatementType.NOTE)
        targets = _build_word_comment_targets(
            [table],
            [TableCheckResult(1, "Có kiểm tra")],
            {1: Workbook().active},
            [
                NoteMatchResult(
                    note_table_index=1,
                    note_title="Chi phí thuế thu nhập doanh nghiệp hiện hành",
                    note_current_cell="B4",
                    note_prior_cell="B4",
                    source="PL",
                    item_name="Tổng lợi nhuận kế toán trước thuế",
                    statement_code="50",
                    current_difference=Decimal(0),
                    prior_difference=Decimal("25"),
                    status="Difference",
                    match_type="Tên dòng + PL",
                )
            ],
        )

        self.assertEqual(len(targets), 1)
        self.assertIn("chênh lệch 25 VND", targets[0].message)

    def test_maturity_match_does_not_create_word_comment(self):
        table = ExtractedTable(
            1,
            [["Ngắn hạn", "100", "90"]],
            StatementType.NOTE,
            "Phải thu khác",
        )
        worksheet = Workbook().active
        targets = _build_word_comment_targets(
            [table],
            [TableCheckResult(1, "Có kiểm tra")],
            {1: worksheet},
            [
                NoteMatchResult(
                    note_table_index=1,
                    note_title=table.title_hint,
                    note_current_cell="B5",
                    note_prior_cell="C5",
                    item_name="Phải thu ngắn hạn khác",
                    statement_code="136",
                    status="Matched",
                    match_type="Tên dòng và Thuyết minh",
                )
            ],
        )

        self.assertEqual(targets, [])

    def test_maturity_title_comment_reports_period_difference_in_vnd(self):
        table = ExtractedTable(1, [["Tổng cộng", "1.100"]], StatementType.NOTE)
        worksheet = Workbook().active
        targets = _build_word_comment_targets(
            [table],
            [TableCheckResult(1, "Có kiểm tra")],
            {1: worksheet},
            [
                NoteMatchResult(
                    note_table_index=1,
                    note_title="15.1. Ngắn hạn",
                    note_current_cell="B5",
                    source="BS",
                    item_name="Vay và nợ thuê tài chính ngắn hạn",
                    statement_code="320",
                    current_difference=Decimal("1000000"),
                    status="Difference",
                    match_type="Tiêu đề kỳ hạn và Thuyết minh",
                )
            ],
        )

        self.assertEqual(
            targets[0].message,
            "Đối chiếu TM–BS: Tìm thấy tại “Vay và nợ thuê tài chính ngắn hạn”, "
            "mã số 320. Có chênh lệch 1.000.000 VND.",
        )

    def test_cash_flow_depreciation_difference_has_detailed_comment(self):
        table = ExtractedTable(
            1,
            [["Chi phí khấu hao tài sản", "110", "90"]],
            StatementType.NOTE,
            "Chi phí sản xuất, kinh doanh theo yếu tố",
        )
        targets = _build_word_comment_targets(
            [table],
            [TableCheckResult(1, "Có kiểm tra")],
            {1: Workbook().active},
            [
                NoteMatchResult(
                    note_table_index=1,
                    note_title=table.title_hint,
                    note_current_cell="B4",
                    note_prior_cell="C4",
                    source="CF",
                    item_name="Khấu hao tài sản cố định và bất động sản đầu tư",
                    statement_code="02",
                    current_difference=Decimal("10"),
                    prior_difference=Decimal("0"),
                    status="Difference",
                    match_type="Tên dòng + mã CF",
                )
            ],
        )

        self.assertEqual(len(targets), 1)
        self.assertEqual(
            targets[0].message,
            "Đối chiếu TM–CF: Tìm thấy tại “Khấu hao tài sản cố định và bất động sản đầu tư”, "
            "mã số 02. Có chênh lệch 10 VND.",
        )
        self.assertNotIn("Số liệu khớp", targets[0].message)

    def test_fractional_difference_comment_keeps_two_decimal_precision(self):
        table = ExtractedTable(
            1,
            [["Chỉ tiêu USD", "100.01"]],
            StatementType.NOTE,
            "Thuyết minh USD",
        )
        targets = _build_word_comment_targets(
            [table],
            [TableCheckResult(1, "Có kiểm tra")],
            {1: Workbook().active},
            [
                NoteMatchResult(
                    note_table_index=1,
                    note_title=table.title_hint,
                    note_current_cell="B5",
                    source="BS",
                    item_name="Chỉ tiêu USD",
                    statement_code="100",
                    current_difference=Decimal("0.01"),
                    status="Difference",
                    match_type="Tên dòng + mã BS",
                )
            ],
        )

        self.assertEqual(
            targets[0].message,
            "Đối chiếu TM–BS: Tìm thấy tại “Chỉ tiêu USD”, mã số 100. "
            "Có chênh lệch 0,01 VND.",
        )

    def test_matched_full_item_title_does_not_create_word_comment(self):
        table = ExtractedTable(
            1,
            [["Tổng cộng", "1.000", "900"]],
            StatementType.NOTE,
            "Phải thu ngắn hạn của khách hàng",
        )
        targets = _build_word_comment_targets(
            [table],
            [TableCheckResult(1, "Có kiểm tra")],
            {1: Workbook().active},
            [
                NoteMatchResult(
                    note_table_index=1,
                    note_title=table.title_hint,
                    statement_code="131",
                    status="Matched",
                    match_type="Tham chiếu Thuyết minh",
                )
            ],
        )

        self.assertEqual(targets, [])

    def test_matched_total_value_fallback_does_not_create_word_comment(self):
        table = ExtractedTable(1, [["Tổng cộng", "1.000"]], StatementType.NOTE)
        targets = _build_word_comment_targets(
            [table],
            [TableCheckResult(1, "Có kiểm tra")],
            {1: Workbook().active},
            [
                NoteMatchResult(
                    note_table_index=1,
                    note_title="Vốn chủ sở hữu",
                    source="BS",
                    item_name="Vốn đầu tư của chủ sở hữu",
                    statement_code="411",
                    status="Matched",
                    match_type="Mã Thuyết minh + tên + Tổng cộng",
                )
            ],
        )

        self.assertEqual(targets, [])

    def test_skips_format_warning_and_uses_short_arithmetic_message(self):
        table = ExtractedTable(1, [["1.000", "100"]], StatementType.NOTE)
        worksheet = Workbook().active
        format_cell = worksheet.cell(DATA_START_ROW, 1, "1.000")
        format_cell.fill = WARNING_FILL
        format_cell.comment = Comment("Định dạng số khác đa số trong bảng.", "CHECK_FS")
        arithmetic_cell = worksheet.cell(DATA_START_ROW, 2, 100)
        arithmetic_cell.fill = WARNING_FILL

        targets = _build_word_comment_targets(
            [table],
            [TableCheckResult(1, "Có kiểm tra", issue_count=2)],
            {1: worksheet},
            [],
        )

        self.assertEqual(len(targets), 1)
        self.assertEqual(targets[0].column_index, 1)
        self.assertEqual(targets[0].message, "Kiểm tra số học: Có chênh lệch.")

    def test_note_reconciliation_message_is_short_and_distinct(self):
        table = ExtractedTable(1, [["Doanh thu", "100"]], StatementType.NOTE)
        worksheet = Workbook().active
        targets = _build_word_comment_targets(
            [table],
            [TableCheckResult(1, "Có kiểm tra")],
            {1: worksheet},
            [
                NoteMatchResult(
                    note_table_index=1,
                    note_title="Doanh thu",
                    note_current_cell="B5",
                    status="Difference",
                )
            ],
        )

        self.assertEqual(len(targets), 1)
        self.assertEqual(
            targets[0].message,
            "Đối chiếu TM–BCTC (Cần xem xét): Phát hiện số liệu nghi vấn chưa khớp giữa TM và BCTC. Đề nghị kiểm toán viên kiểm tra lại cấu trúc Thuyết minh.",
        )

    def test_standard_note_reconciliation_comment_includes_difference_amount(self):
        table = ExtractedTable(
            1,
            [["Tổng cộng", "16.413.246.268", "2.871.741.001"]],
            StatementType.NOTE,
            "Tiền",
        )
        targets = _build_word_comment_targets(
            [table],
            [TableCheckResult(1, "Có kiểm tra")],
            {1: Workbook().active},
            [
                NoteMatchResult(
                    note_table_index=1,
                    note_title="Tiền",
                    note_current_cell="B4",
                    current_difference=Decimal("8"),
                    status="Difference",
                    match_type="Tham chiếu Thuyết minh",
                )
            ],
        )

        self.assertEqual(len(targets), 1)
        self.assertEqual(
            targets[0].message,
            "Đối chiếu TM–BCTC (Cần xem xét): Phát hiện số liệu nghi vấn chưa khớp giữa TM và BCTC (chênh lệch 8 VND). Đề nghị kiểm toán viên kiểm tra lại cấu trúc Thuyết minh.",
        )

    def test_skips_unlinked_note_without_statement_target(self):
        table = ExtractedTable(
            1,
            [["Chi phí nguyên liệu", "100"], ["Tổng cộng", "100"]],
            StatementType.NOTE,
            "Chi phí sản xuất, kinh doanh theo yếu tố",
        )
        worksheet = Workbook().active
        targets = _build_word_comment_targets(
            [table],
            [TableCheckResult(1, "Có kiểm tra")],
            {1: worksheet},
            [NoteMatchResult(note_table_index=1, note_title=table.title_hint, status="Statement not found")],
        )
        self.assertEqual(targets, [])

    def test_keeps_missing_statement_comment_when_target_code_is_known(self):
        table = ExtractedTable(1, [["Nguyên giá", "100"]], StatementType.NOTE)
        worksheet = Workbook().active
        targets = _build_word_comment_targets(
            [table],
            [TableCheckResult(1, "Có kiểm tra")],
            {1: worksheet},
            [NoteMatchResult(note_table_index=1, note_title="Tài sản cố định hữu hình", statement_code="222", status="Statement not found")],
        )
        self.assertEqual(len(targets), 1)
        self.assertEqual(targets[0].message, "Đối chiếu TM–BCTC: Không tìm thấy chỉ tiêu BCTC.")

    def test_zero_missing_bs_item_uses_one_comment_for_both_periods(self):
        table = ExtractedTable(
            1,
            [["Thuế và các khoản khác phải thu Nhà nước", "-", "", "-"]],
            StatementType.NOTE,
        )
        targets = _build_word_comment_targets(
            [table],
            [TableCheckResult(1, "Có kiểm tra")],
            {1: Workbook().active},
            [
                NoteMatchResult(
                    note_table_index=1,
                    note_title="Thuế và các khoản phải nộp Nhà nước",
                    note_current=Decimal(0),
                    note_current_cell="B4",
                    note_prior=Decimal(0),
                    note_prior_cell="D4",
                    statement_code="153",
                    item_name="Thuế và các khoản khác phải thu Nhà nước",
                    status="Statement not found",
                    match_type="Tên dòng + mã BS",
                )
            ],
        )

        self.assertEqual(len(targets), 1)
        self.assertEqual((targets[0].row_index, targets[0].column_index), (0, 1))

    def test_named_total_without_statement_target_creates_review_comments(self):
        table = ExtractedTable(
            1,
            [["", "Số cuối năm", "Số đầu năm"], ["Tổng cộng", "273", "35"]],
            StatementType.NOTE,
            "Phải thu khác",
        )
        targets = _build_word_comment_targets(
            [table],
            [TableCheckResult(1, "Có kiểm tra")],
            {1: Workbook().active},
            [
                NoteMatchResult(
                    note_table_index=1,
                    note_title=table.title_hint,
                    note_current_cell="B5",
                    note_prior_cell="C5",
                    item_name="Phải thu khác",
                    status="Statement not found",
                    match_type="Tên bảng + Tổng cộng",
                )
            ],
        )

        self.assertEqual(len(targets), 2)
        self.assertTrue(all("Chưa tìm thấy chỉ tiêu phù hợp" in target.message for target in targets))


if __name__ == "__main__":
    unittest.main()
